"""Traceable plant-layout study reports from the checked canonical dataset.

HTML and tabular evidence are always written.  PDF and Word authoring are
optional adapters; their availability is explicit in the returned manifest.
No report state is presented as engineering design approval.
"""

from __future__ import annotations

import csv
import html
import json
from collections import Counter
from pathlib import Path

from layout_model import equipment_bounds

DISCLAIMER = (
    "This concept layout study compares preliminary equipment arrangements. "
    "Equipment dimensions and rules retain their recorded source and uncertainty. "
    "Open hazard, structural, vendor, lifting and evacuation requirements require "
    "specialist evidence before design acceptance. No alternative is approved by this report."
)


def _value(value):
    """Format finite JSON values for readable report tables."""
    if value is None:
        return "unavailable"
    if isinstance(value, float):
        return "%.3f" % value
    if isinstance(value, (dict, list)):
        return json.dumps(value, sort_keys=True)
    return str(value)


def _counts(alternative):
    """Count every explicit check status, including non-applicable requirements."""
    c = Counter(x["status"] for x in alternative.get("checks", []))
    return [c[s] for s in ("PASS", "FAIL", "NOT_EVALUATED", "NOT_APPLICABLE")]


def _comparison(alternatives):
    """Produce a quantitative comparison without ranking incomplete designs."""
    headers = ["Alternative", "Status", "Pipe proxy m", "Area m2", "Mass kg", "Pass", "Fail", "Open", "N/A"]
    rows = []
    for a in alternatives:
        m = a.get("metrics", {})
        rows.append([a["id"], a.get("status", "NOT_EVALUATED"),
                     _value(m.get("piping_proxy_m")), _value(m.get("occupied_area_m2")),
                     _value(m.get("total_mass_kg"))] + _counts(a))
    return headers, rows


def _coordinates(study, alternatives):
    """Return the exact placement and physical envelope coordinate schedule."""
    equipment = {e["tag"]: e for e in study["equipment"]}
    decks = {d["id"]: d for d in study["decks"]}
    headers = ["alternative", "tag", "deck", "x_m", "y_m", "z_m", "rotation_deg",
               "length_m", "width_m", "height_m", "mass_kg", "dimension_source"]
    rows = []
    for a in alternatives:
        for p in sorted(a.get("placements", []), key=lambda x: x["tag"]):
            e, d = equipment[p["tag"]], decks[p["deck"]]
            b = equipment_bounds(e, p, d)
            rows.append([a["id"], p["tag"], p["deck"], p["x_m"], p["y_m"], b[2],
                         p.get("rotation_deg", 0)] + list(e["size_m"]) +
                        [e.get("mass_kg"), e.get("dimension_source", {}).get("reference", "unavailable")])
    return headers, rows


def _table_html(headers, rows):
    """Escape all model text before inserting an HTML evidence table."""
    return "<table><thead><tr>" + "".join("<th>%s</th>" % html.escape(str(v)) for v in headers) + \
        "</tr></thead><tbody>" + "".join("<tr>" + "".join("<td>%s</td>" % html.escape(_value(v))
        for v in row) + "</tr>" for row in rows) + "</tbody></table>"


def _drawing_groups(drawings, alternatives):
    """Group actual generated plan sheets by alternative in controlled order."""
    return {a["id"]: [s for s in (drawings or {}).get("sheets", []) if s.get("alternative") == a["id"]]
            for a in alternatives}


def _source_rows(study):
    """List source metadata without implying that catalogue access proves clauses."""
    return [[s.get("document", "unassigned"), s.get("edition", "unassigned"),
             s.get("access_status", "unavailable"), s.get("reference", "unavailable")]
            for s in study.get("standards", [])]


def _open_checks(alternative):
    """Return all failures and unevaluated checks for the action register."""
    return [c for c in alternative.get("checks", []) if c["status"] in ("FAIL", "NOT_EVALUATED")]


def _process_evidence(study):
    """Return process snapshot evidence without inferring physical dimensions."""
    imported = study.get("process_import")
    if not imported:
        return None
    data = imported.get("source_data", {})
    meta = data.get("meta", {})
    summary = [["Execution status", imported.get("execution_status", "SOURCE_EXTRACT_ONLY")],
               ["Source", imported.get("source", imported.get("reference", meta.get("model_name", "recorded snapshot")))],
               ["NeqSim version", meta.get("neqsim_version", "unavailable")],
               ["EOS", meta.get("eos", "unavailable")],
               ["Snapshot fingerprint", imported.get("fingerprint", "unavailable")]]
    streams, equipment = [], []
    for stream in data.get("streams", []):
        for case, values in stream.get("results", {}).items():
            streams.append([stream.get("id", "unassigned"), case, values.get("t_c"),
                            values.get("p_bara"), values.get("mass_kg_h")])
    for unit in data.get("equipment", []):
        for case, values in unit.get("results", {}).items():
            if values.get("power_kw") is not None or values.get("duty_kw") is not None:
                equipment.append([unit.get("tag", "unassigned"), case,
                                  values.get("power_kw"), values.get("duty_kw")])
    return {"summary": summary, "streams": streams, "equipment": equipment}


def _cad_summary(cad):
    """Summarize adapters and exchange checks without dumping scratch paths."""
    if not cad:
        return [["CAD export", "Not requested"]]
    rows = []
    for name, info in cad.get("adapters", {}).items():
        rows.append([name, "%s; version %s" % (info.get("status", "unavailable"), info.get("version", "unavailable"))])
    formats = Counter(item.get("format", "unknown") for item in cad.get("files", []))
    rows.extend([[str(fmt).upper() + " files", n] for fmt, n in sorted(formats.items())])
    counts = Counter(item.get("status", "NOT_EVALUATED") for item in cad.get("checks", []))
    rows.append(["Exchange checks", ", ".join("%s %d" % (key, val) for key, val in sorted(counts.items()))])
    rows.extend([["Warning", warning] for warning in cad.get("warnings", [])])
    return rows


def _inventory_tables(inventory):
    """Return compact physical-route, isolation-section and blocker tables."""
    inventory = inventory or {}
    routes = [[route["id"], route["connection_id"], route["line_tag"], route["length_m"],
               route["internal_diameter_m"], route["pipe_volume_m3"],
               route["source"]["reference"], route["source"]["revision"],
               route["source"]["evidence_status"]]
              for route in inventory.get("routes", [])]
    sections = [[section["id"], section["route_ids"], section["equipment_volume_ids"],
                 section["boundary_device_ids"], section["pipe_volume_m3"],
                 section["equipment_volume_m3"], section["total_isolated_volume_m3"]]
                for section in inventory.get("isolation_sections", [])]
    blockers = [[blocker.get("code"), blocker.get("subject", "study"), blocker.get("message")]
                for blocker in inventory.get("blockers", [])]
    return routes, sections, blockers


def _html_report(study, alternatives, drawings, cad, inventory=None):
    """Build a portable HTML report with inline vector drawing views."""
    esc = lambda v: html.escape(_value(v))
    p = study["project"]
    headers, rows = _comparison(alternatives)
    coordinate_headers, coordinate_rows = _coordinates(study, alternatives)
    route_rows, section_rows, blocker_rows = _inventory_tables(inventory)
    title = p.get("title", "Plant layout study")
    out = ["<!doctype html><html lang='en'><meta charset='utf-8'>",
           "<meta name='viewport' content='width=device-width,initial-scale=1'>",
           "<title>%s</title>" % esc(title),
           "<style>body{font:16px/1.5 system-ui,sans-serif;max-width:1180px;margin:36px auto;padding:0 24px;color:#1b2938}"
           "h1{font-size:34px;line-height:1.2}h2{margin-top:36px}h3{margin-top:28px}"
           "table{border-collapse:collapse;width:100%;font-size:13px;margin:16px 0;overflow-wrap:anywhere}"
           "th,td{text-align:left;padding:8px;border:1px solid #cbd4dc;vertical-align:top}"
           "th{background:#edf2f6}small{color:#526372}.notice{background:#fff3dc;border-left:4px solid #b87922;padding:16px}"
           ".drawing{overflow:auto;border:1px solid #d8dfe5;margin:18px 0}.drawing svg{width:100%;height:auto;display:block}"
           "code{overflow-wrap:anywhere}a{color:#174e7c}@media print{body{margin:0}h2{break-after:avoid}tr{break-inside:avoid}}"
           "</style><body>", "<small>%s | Revision %s | %s</small>" % (esc(p["id"]), esc(p.get("revision", "A")),
                                                                                         esc(p.get("maturity", "SCREENING"))),
           "<h1>%s</h1>" % esc(title), "<p class='notice'>%s</p>" % esc(DISCLAIMER),
           "<p>The study contains %d equipment items, %d plot or deck levels, and %d generated alternatives. "
           "Their geometry, quantitative comparison and unresolved requirements are reported together.</p>" %
           (len(study["equipment"]), len(study["decks"]), len(alternatives)),
           "<h2>Alternative comparison</h2>", _table_html(headers, rows),
           "<p>Pipe proxy is a centre-to-centre distance metric for arrangement comparison. "
           "It is not a routed piping quantity or a pressure-drop calculation. Feasibility is evaluated "
           "against modelled constraints; open required evidence keeps an alternative incomplete.</p>",
           "<h2>Design basis</h2>"]
    out.append(_table_html(["Basis", "Value"], [["Facility", p.get("facility_type")],
               ["Jurisdiction", p.get("jurisdiction")], ["Coordinate system", study.get("coordinate_system")],
               ["Placement options", study.get("options")]]))
    for key, label in (("assumptions", "Assumptions"), ("data_gaps", "Data gaps")):
        out.append("<h3>%s</h3><ul>" % label)
        out.extend("<li>%s</li>" % esc(v) for v in study.get(key, []))
        if not study.get(key):
            out.append("<li>No items recorded.</li>")
        out.append("</ul>")
    process = _process_evidence(study)
    if process:
        out.extend(["<h2>Process model evidence</h2>",
                    "<p>These results belong to the imported process snapshot. Installed package geometry retains "
                    "its independent dimension source; stream results do not establish vendor footprints.</p>",
                    _table_html(["Item", "Value"], process["summary"]),
                    _table_html(["Stream", "Case", "Temperature C", "Pressure bara", "Mass flow kg/h"], process["streams"]),
                    _table_html(["Equipment", "Case", "Shaft power kW", "Heat duty kW"], process["equipment"])])
    out.extend(["<h2>Physical piping and isolation inventory</h2>",
                "<p>Only supplied routed centrelines, internal diameters, equipment volumes and isolation "
                "boundaries are included. Process connectivity and centre-to-centre proxies are never used "
                "as pipe length or isolated inventory.</p>",
                _table_html(["Status", "Ready for dynamic safeguarding design", "Calculation basis"],
                            [[(inventory or {}).get("status", "NOT_EVALUATED"),
                              (inventory or {}).get("ready_for_dynamic_safeguarding_design", False),
                              (inventory or {}).get("calculation_basis", "No inventory supplied")]]),
                _table_html(["Route", "Connection", "Line", "Length m", "ID m", "Volume m3",
                             "Source", "Revision", "Evidence status"], route_rows),
                _table_html(["Isolation section", "Routes", "Equipment volumes", "Boundary devices",
                             "Pipe m3", "Equipment m3", "Total isolated m3"], section_rows),
                "<h3>Physical inventory blockers</h3>",
                _table_html(["Code", "Subject", "Blocker"], blocker_rows)])
    groups = _drawing_groups(drawings, alternatives)
    for a in alternatives:
        out.append("<h2>%s</h2><p>%s</p>" % (esc(a["id"]), esc(a.get("label", ""))))
        out.append(_table_html(["Solver field", "Value"], [[k, v] for k, v in a.get("solver", {}).items()]))
        for sheet in groups[a["id"]]:
            path = Path(sheet["svg"])
            if path.is_file():
                svg = path.read_text(encoding="utf-8")
                svg = svg[svg.find("<svg"):]
                out.append("<h3>%s</h3><div class='drawing'>%s</div>" % (esc(sheet["title"]), svg))
        findings = _open_checks(a)
        out.append("<h3>Findings requiring action</h3>")
        out.append(_table_html(["Check", "Status", "Affected equipment", "Finding", "Source"],
                   [[c["id"], c["status"], ", ".join(c.get("equipment", [])), c.get("message", ""), c.get("source", {})]
                    for c in findings]))
        if not findings:
            out.append("<p>No failed or unevaluated checks are recorded within the implemented scope.</p>")
    out.extend(["<h2>Equipment coordinate schedule</h2>", _table_html(coordinate_headers, coordinate_rows),
                "<h2>Standards and source register</h2>",
                "<p>Access status is separate from applicability and from clause verification. "
                "Metadata-only entries are references for further engineering review.</p>",
                _table_html(["Document", "Edition", "Access status", "Reference"], _source_rows(study)),
                "<h2>Exchange coverage and limitations</h2>",
                "<p>All physical checks concern the documented conceptual rectangular envelope model. "
                "Vendor internals, nozzle orientation, route generation, pipe sizing/hydraulics/stress, structural adequacy, fire and blast loads, "
                "ventilation, human performance and evacuation under accident scenarios require appropriate specialist evidence.</p>",
                _table_html(["Adapter or exchange evidence", "Result"], _cad_summary(cad)),
                "<h2>Complete check evidence</h2>"])
    for a in alternatives:
        out.append("<h3>%s</h3>" % esc(a["id"]))
        out.append(_table_html(["ID", "Kind", "Status", "Required", "Measured", "Limit", "Unit", "Message"],
                   [[c["id"], c.get("kind"), c["status"], c.get("required"), c.get("measured"),
                     c.get("limit"), c.get("unit"), c.get("message", "")] for c in a.get("checks", [])]))
    out.append("</body></html>\n")
    return "\n".join(out)


def _pdf_report(path, study, alternatives, drawings, cad=None, inventory=None):
    """Generate a paginated report with wrapped evidence tables and plan views."""
    import matplotlib
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (Image, KeepTogether, PageBreak, Paragraph,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    # Embed the Matplotlib-bundled, openly licensed DejaVu fonts.  Depending on
    # a PDF viewer's Base14 Helvetica substitute can change glyph advances and
    # produce unreadable spacing in headless or font-restricted environments.
    fonts = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
    for name, filename in (("LayoutSans", "DejaVuSans.ttf"),
                           ("LayoutSans-Bold", "DejaVuSans-Bold.ttf"),
                           ("LayoutSans-Oblique", "DejaVuSans-Oblique.ttf"),
                           ("LayoutSans-BoldOblique", "DejaVuSans-BoldOblique.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(fonts / filename)))
    pdfmetrics.registerFontFamily("LayoutSans", normal="LayoutSans", bold="LayoutSans-Bold",
                                  italic="LayoutSans-Oblique", boldItalic="LayoutSans-BoldOblique")

    styles = getSampleStyleSheet()
    styles["BodyText"].fontName = "LayoutSans"
    styles["Title"].fontName = "LayoutSans-Bold"
    styles["Heading1"].fontName = "LayoutSans-Bold"
    styles["Heading2"].fontName = "LayoutSans-Bold"
    styles.add(ParagraphStyle("LayoutBody", parent=styles["BodyText"], fontName="LayoutSans", fontSize=9,
                              leading=13, spaceAfter=6))
    styles.add(ParagraphStyle("LayoutSmall", parent=styles["BodyText"], fontSize=7.5, leading=10))
    styles["Title"].fontSize = 23
    styles["Title"].leading = 28
    styles["Title"].alignment = TA_LEFT
    styles["Title"].textColor = colors.HexColor("#172b3f")
    styles["Heading1"].fontSize = 14
    styles["Heading1"].leading = 18
    styles["Heading1"].spaceBefore = 12
    styles["Heading1"].spaceAfter = 8
    styles["Heading1"].keepWithNext = True
    styles["Heading2"].keepWithNext = True
    body = []

    def paragraph(text, style="LayoutBody"):
        return Paragraph(html.escape(_value(text)).replace("\n", "<br/>"), styles[style])

    def table(headers, rows, widths):
        cells = [[paragraph(v, "LayoutSmall") for v in headers]] + \
            [[paragraph(v, "LayoutSmall") for v in row] for row in rows]
        tab = Table(cells, colWidths=[w * mm for w in widths], repeatRows=1, hAlign="LEFT")
        tab.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef3")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#bac6cf")),
            ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d9df")),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        return tab

    p = study["project"]
    body.extend([paragraph(p.get("title", "Plant layout study"), "Title"),
                 paragraph("%s   Revision %s   SCREENING" % (p["id"], p.get("revision", "A"))),
                 Spacer(1, 3 * mm), paragraph(DISCLAIMER),
                 paragraph("The study compares %d alternatives for %d equipment items across %d plot or deck levels. "
                           "It retains explicit model checks, source metadata and unresolved specialist requirements." %
                           (len(alternatives), len(study["equipment"]), len(study["decks"]))),
                 paragraph("Alternative comparison", "Heading1")])
    headers, rows = _comparison(alternatives)
    # Compact comparison includes the critical readiness measures; the HTML/CSV keeps all metrics.
    body.append(table([headers[i] for i in (0, 1, 2, 3, 6, 7)],
                      [[r[i] for i in (0, 1, 2, 3, 6, 7)] for r in rows], [23, 33, 28, 25, 25, 30]))
    body.append(paragraph("Pipe proxy is a centre-to-centre arrangement metric, not detailed routing. "
                          "A low proxy does not compensate for failed or unevaluated requirements."))
    body.append(paragraph("Design basis", "Heading1"))
    basis_rows = [["Facility", p.get("facility_type")], ["Jurisdiction", p.get("jurisdiction")],
                  ["Length and mass units", "m and kg"],
                  ["Elevation datum", study.get("coordinate_system", {}).get("elevation_datum", "project")],
                  ["Equipment dimensions", "Source status retained for each equipment item"],
                  ["Engineering approval", "Not assigned"]]
    body.append(table(["Basis", "Value"], basis_rows, [45, 119]))
    for key, title in (("assumptions", "Assumptions"), ("data_gaps", "Data gaps")):
        body.append(paragraph(title, "Heading1"))
        for item in study.get(key, []) or ["No items recorded."]:
            body.append(paragraph("- " + str(item)))
    process = _process_evidence(study)
    if process:
        body.append(paragraph("Process model evidence", "Heading1"))
        body.append(paragraph("The imported snapshot supplies the following process conditions. Package geometry "
                              "retains its independent dimension source and is not inferred from stream results."))
        body.append(table(["Item", "Value"], process["summary"], [43, 121]))
        body.append(table(["Stream", "Case", "T C", "P bara", "Flow kg/h"], process["streams"], [45, 22, 30, 30, 37]))
        if process["equipment"]:
            body.append(table(["Equipment", "Case", "Power kW", "Duty kW"], process["equipment"], [55, 25, 42, 42]))
    route_rows, section_rows, blocker_rows = _inventory_tables(inventory)
    body.append(PageBreak())
    body.append(paragraph("Physical piping and isolation inventory", "Heading1"))
    body.append(paragraph("Status: %s. Only supplied routed geometry and volumes are included; process connectivity "
                          "and arrangement proxies are not substituted for actual pipe length or inventory." %
                          (inventory or {}).get("status", "NOT_EVALUATED")))
    totals = (inventory or {}).get("totals", {})
    body.append(table(["Quantity", "Value", "Unit"],
                      [["Routed centreline", totals.get("route_length_m"), "m"],
                       ["Pipe inventory", totals.get("pipe_volume_m3"), "m3"],
                       ["Equipment inventory", totals.get("equipment_volume_m3"), "m3"]],
                      [65, 55, 44]))
    if route_rows:
        body.append(table(["Route", "Connection", "Length m", "ID m", "Volume m3"],
                          [[row[i] for i in (0, 1, 3, 4, 5)] for row in route_rows], [35, 42, 28, 27, 32]))
    if section_rows:
        body.append(table(["Section", "Routes", "Boundary devices", "Total m3"],
                          [[row[i] for i in (0, 1, 3, 6)] for row in section_rows], [35, 43, 54, 32]))
    if blocker_rows:
        body.append(table(["Blocker", "Subject", "Action"], blocker_rows, [43, 38, 83]))
    else:
        body.append(paragraph("No physical-inventory completeness blockers are recorded for the declared scope."))
    groups = _drawing_groups(drawings, alternatives)
    for a in alternatives:
        solver = a.get("solver", {})
        grid = solver.get("settings", {}).get("grid_m", "unavailable")
        body.extend([PageBreak(), paragraph(a["id"] + " arrangement", "Heading1"),
                     paragraph(a.get("label", "")),
                     paragraph("Status: %s. %s placement search returned %s on a %s m grid after %s candidate evaluations. "
                               "Global optimality is not established." %
                               (a.get("status", "NOT_EVALUATED"), solver.get("backend", "unavailable"),
                                solver.get("status", "unavailable"), grid,
                                solver.get("candidates_evaluated", "unavailable")))])
        sheets = groups[a["id"]]
        # Report previews support the narrative; the separate A3 package is the drawing deliverable.
        for s in sheets:
            image = Path(s.get("png", ""))
            if image.is_file():
                body.append(KeepTogether([paragraph(s["title"], "Heading2"),
                                          Image(str(image), width=164 * mm, height=116 * mm),
                                          paragraph("Reduced A3 preview. Read dimensions in the full-size drawing package.", "LayoutSmall")]))
        findings = _open_checks(a)
        body.append(paragraph("Findings requiring action", "Heading1"))
        if findings:
            body.append(table(["Check", "Status", "Finding and affected equipment"],
                       [[c["id"], c["status"], c.get("message", "") +
                         (" Equipment: " + ", ".join(c.get("equipment", [])) if c.get("equipment") else "")]
                        for c in findings], [35, 30, 99]))
        else:
            body.append(paragraph("No failed or unevaluated checks are recorded within the implemented scope."))
    body.extend([PageBreak(), paragraph("Equipment coordinate schedule", "Heading1")])
    coord_h, coord_rows = _coordinates(study, alternatives)
    body.append(table(["Alt", "Tag", "Deck", "E m", "N m", "EL m", "Rot deg"],
                      [[r[i] for i in range(7)] for r in coord_rows], [19, 35, 23, 23, 23, 23, 18]))
    body.append(paragraph("The full CSV adds source dimensions, mass and dimension references. "
                          "Equipment orientation uses the local vertical axis; shape sizes are unrotated source dimensions."))
    body.append(paragraph("Standards and source register", "Heading1"))
    source_rows = _source_rows(study)
    if source_rows:
        body.append(table(["Document", "Edition", "Access status"], [r[:3] for r in source_rows], [64, 45, 55]))
    body.append(paragraph("Public catalogue metadata establishes document identity and scope, not detailed clause compliance. "
                          "The HTML report and canonical dataset retain source links and every calculated check."))
    body.append(paragraph("Limitations and next engineering work", "Heading1"))
    body.append(paragraph("The model checks rectangular conceptual envelopes and reserved corridors. "
                          "Vendor package confirmation, nozzles, route generation, pipe sizing/hydraulics/stress, fire and blast, "
                          "dispersion and ventilation, structural loads, lifting operations and evacuation under "
                          "accident scenarios require appropriately scoped engineering evidence. "
                          "Existing model checks must be rerun after layout or basis changes."))
    body.append(paragraph("CAD exchange evidence", "Heading1"))
    body.append(table(["Adapter or exchange evidence", "Result"], _cad_summary(cad), [65, 99]))
    body.append(paragraph("Traceability", "Heading1"))
    body.append(paragraph("Read dataset/plant_layout.json, reports/coordinate_schedule.csv, "
                          "reports/check_results.json and drawings/drawing_spec.json for machine-readable results. "
                          "These files share equipment tags and alternative IDs with the drawings."))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("LayoutSans", 8)
        canvas.setFillColor(colors.HexColor("#566675"))
        canvas.drawString(23 * mm, 13 * mm, p["id"] + " | SCREENING | Rev " + str(p.get("revision", "A")))
        canvas.drawRightString(187 * mm, 13 * mm, "Page %d" % doc.page)
        canvas.restoreState()

    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=23 * mm, leftMargin=23 * mm,
                            topMargin=20 * mm, bottomMargin=22 * mm,
                            title=p.get("title", "Plant layout study"), author="NeqSim layout tooling")
    doc.build(body, onFirstPage=footer, onLaterPages=footer)


def _docx_report(path, study, alternatives, drawings, cad=None, inventory=None):
    """Write an editable Word companion with controlled section and table styles."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.name = "Arial"
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        properties = doc.styles[name].element.find(qn("w:pPr"))
        if properties is not None:
            for border in properties.findall(qn("w:pBdr")):
                properties.remove(border)
    p = study["project"]
    doc.add_paragraph(p.get("title", "Plant layout study"), style="Title")
    doc.add_paragraph("%s   Revision %s   SCREENING" % (p["id"], p.get("revision", "A")))
    doc.add_paragraph(DISCLAIMER)

    def table(headers, rows):
        tab = doc.add_table(rows=1, cols=len(headers))
        tab.style = "Light Shading Accent 1"
        for cell, text in zip(tab.rows[0].cells, headers):
            cell.text = str(text)
        repeat = OxmlElement("w:tblHeader")
        tab.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for cell, text in zip(tab.add_row().cells, row):
                cell.text = _value(text)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(8)
        return tab

    doc.add_heading("Alternative comparison", 1)
    headers, rows = _comparison(alternatives)
    table([headers[i] for i in (0, 1, 2, 6, 7)], [[r[i] for i in (0, 1, 2, 6, 7)] for r in rows])
    doc.add_paragraph("Pipe proxy is a centre-to-centre arrangement metric and does not establish routed piping quantities.")
    doc.add_heading("Design basis", 1)
    table(["Item", "Value"], [["Facility", p.get("facility_type")], ["Jurisdiction", p.get("jurisdiction")],
                               ["Coordinate units", "metres; +Z up"], ["Mass units", "kg"]])
    for key, heading in (("assumptions", "Assumptions"), ("data_gaps", "Data gaps")):
        doc.add_heading(heading, 1)
        for item in study.get(key, []) or ["No items recorded."]:
            doc.add_paragraph(str(item), style="List Bullet")
    process = _process_evidence(study)
    if process:
        doc.add_heading("Process model evidence", 1)
        doc.add_paragraph("Process results belong to the imported snapshot. Package footprints retain their independent dimension provenance.")
        table(["Item", "Value"], process["summary"])
        table(["Stream", "Case", "T C", "P bara", "Flow kg/h"], process["streams"])
        if process["equipment"]:
            table(["Equipment", "Case", "Power kW", "Duty kW"], process["equipment"])
    route_rows, section_rows, blocker_rows = _inventory_tables(inventory)
    doc.add_heading("Physical piping and isolation inventory", 1)
    doc.add_paragraph("Status: %s. Only supplied route geometry and contained volumes are used." %
                      (inventory or {}).get("status", "NOT_EVALUATED"))
    table(["Route", "Connection", "Length m", "ID m", "Volume m3"],
          [[row[i] for i in (0, 1, 3, 4, 5)] for row in route_rows])
    table(["Section", "Routes", "Boundary devices", "Total m3"],
          [[row[i] for i in (0, 1, 3, 6)] for row in section_rows])
    if blocker_rows:
        table(["Blocker", "Subject", "Action"], blocker_rows)
    groups = _drawing_groups(drawings, alternatives)
    for a in alternatives:
        doc.add_page_break()
        doc.add_heading(a["id"] + " arrangement", 1)
        doc.add_paragraph("Status: " + a.get("status", "NOT_EVALUATED"))
        doc.add_paragraph(a.get("label", ""))
        for sheet in groups[a["id"]]:
            img = Path(sheet.get("png", ""))
            if img.is_file():
                doc.add_heading(sheet["title"], 2)
                doc.add_picture(str(img), width=Inches(6.6))
                doc.add_paragraph("Reduced A3 preview. Refer to the full-size drawing package.", style="Caption")
        doc.add_heading("Findings requiring action", 2)
        findings = _open_checks(a)
        table(["Check", "Status", "Finding"], [[c["id"], c["status"], c.get("message", "")] for c in findings])
    doc.add_page_break()
    doc.add_heading("Standards and source register", 1)
    table(["Document", "Edition", "Access status"], [r[:3] for r in _source_rows(study)])
    doc.add_paragraph("Detailed references and complete check evidence are retained in the HTML report and canonical dataset. "
                      "Catalogue access is separate from inspected clause evidence and design qualification.")
    doc.add_heading("CAD exchange evidence", 1)
    table(["Adapter or exchange evidence", "Result"], _cad_summary(cad))
    doc.add_heading("Limitations and traceability", 1)
    doc.add_paragraph("The geometry consists of conceptual equipment envelopes. Vendor dimensions, route generation, "
                      "pipe sizing/hydraulics/stress, "
                      "structural adequacy, hazard effects, lifting operations and evacuation assessment require specialist evidence. "
                      "The CSV coordinate schedule and JSON check register use the same tags and alternative IDs as the drawings.")
    doc.core_properties.author = "NeqSim layout tooling"
    doc.core_properties.title = p.get("title", "Plant layout study")
    doc.save(str(path))


def build_report(study, alternatives, out_dir, drawings=None, cad=None, inventory=None):
    """Write the report, exact coordinate CSV and complete JSON check evidence."""
    if inventory is None:
        from layout_inventory import build_inventory
        inventory = build_inventory(study)
    root = Path(out_dir) / "reports"
    root.mkdir(parents=True, exist_ok=True)
    html_path = root / "plant_layout_report.html"
    html_path.write_text(_html_report(study, alternatives, drawings, cad, inventory), encoding="utf-8")
    coord_path = root / "coordinate_schedule.csv"
    headers, rows = _coordinates(study, alternatives)
    with coord_path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(headers)
        writer.writerows(rows)
    checks_path = root / "check_results.json"
    checks_path.write_text(json.dumps({a["id"]: a.get("checks", []) for a in alternatives}, indent=2), encoding="utf-8")
    warnings, files = [], [{"path": str(html_path), "format": "html"},
                           {"path": str(coord_path), "format": "csv"},
                           {"path": str(checks_path), "format": "json"}]
    pdf_path, docx_path = root / "plant_layout_report.pdf", root / "plant_layout_report.docx"
    produced = {}
    for format_, path, builder in (("pdf", pdf_path, _pdf_report), ("docx", docx_path, _docx_report)):
        # These stable output names belong to this generator.  Removing a prior
        # generation prevents an unavailable adapter from exposing stale bytes.
        path.unlink(missing_ok=True)
        try:
            builder(path, study, alternatives, drawings, cad, inventory)
        except ImportError as exc:
            warnings.append("%s report adapter unavailable: %s" % (format_.upper(), exc))
            path.unlink(missing_ok=True)
        else:
            if not path.is_file():
                raise RuntimeError("%s report adapter returned without producing its output" % format_)
            produced[format_] = str(path)
            files.append({"path": str(path), "format": format_})
    return {"html": str(html_path), "pdf": produced.get("pdf"),
            "docx": produced.get("docx"),
            "coordinates": str(coord_path), "checks": str(checks_path),
            "files": files, "warnings": warnings,
            "status": "COMPLETE" if "pdf" in produced else "INCOMPLETE",
            "required_formats": ["html", "pdf"], "optional_formats": ["docx"]}
