"""Validation checks for process-study reports (machine-readable result, fails closed).

``validate_all`` runs data, basis, balance, diagram, export and document checks
and returns a dict with one entry per check (PASS / FAIL / WARN / SKIP). The
function never raises for a failed check; callers decide on the exit code.
"""
import csv
import datetime
import json
import os
import re

import psr_basis
import psr_connectivity
import psr_draft

try:  # report_kit lives in the task template package
    from task_template.step3_report import report_kit
except ImportError:  # pragma: no cover - flat checkout without package install
    report_kit = None

COMP_SUM_TOL = 1.0e-6
STANDARDS_MAX_AGE_DAYS = 90
PHASE_SUM_TOL = 1.0e-5
BAD_TOKENS = re.compile(r"\b(nan|NaN|inf|None)\b")
SECTION_REF_RE = re.compile(r"\b(Sections?|Appendix|Appendices|Tables?|Figures?)\s+[A-Z]?\d+(?:\.\d+)*")


def _check(cid, description, status, detail=""):
    return {"id": cid, "description": description, "status": status, "detail": detail}


def _docx_inventory(path):
    """Return (all_text, tables[(caption, header, rows)], pictures[(descr, caption)], headings)."""
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(path)
    body = doc.element.body
    paragraphs = {p._p: p for p in doc.paragraphs}
    tables = {t._tbl: t for t in doc.tables}
    last_caption = ""
    texts, tabs, pics, heads = [], [], [], []
    pending_pic = None
    for child in body.iterchildren():
        if child in paragraphs:
            p = paragraphs[child]
            txt = p.text
            texts.append(txt)
            style = p.style.name if p.style is not None else ""
            if style.startswith("Heading"):
                heads.append((style, txt))
            if style == "Caption":
                last_caption = txt
                if pending_pic is not None:
                    pics.append((pending_pic, txt))
                    pending_pic = None
            blips = child.findall(".//" + qn("w:drawing"))
            for drawing in blips:
                doc_pr = drawing.find(".//" + qn("wp:docPr"))
                pending_pic = doc_pr.get("descr", "") if doc_pr is not None else ""
        elif child in tables:
            t = tables[child]
            rows = [[c.text for c in r.cells] for r in t.rows]
            tabs.append((last_caption, rows[0] if rows else [], rows[1:]))
            if "Validation checks" not in last_caption:
                texts.extend(" ".join(r) for r in rows)
    if pending_pic is not None:
        pics.append((pending_pic, ""))
    return "\n".join(texts), tabs, pics, heads


def _numbers_from(obj, found):
    if isinstance(obj, bool):
        return
    if isinstance(obj, (int, float)):
        found.add(float(obj))
    elif isinstance(obj, str) and report_kit is not None:
        found.update(report_kit.numbers_in_text(obj, significant_only=False))
    elif isinstance(obj, dict):
        for v in obj.values():
            _numbers_from(v, found)
    elif isinstance(obj, (list, tuple)):
        for v in obj:
            _numbers_from(v, found)


def validate_all(data, basis, standards, balances, diagrams, docx_path=None, facts_path=None, exports_dir=None):
    """Run every check. ``diagrams`` is the merged manifest of ``psr_diagrams`` (or None)."""
    checks = []
    cases = [c["id"] for c in data["cases"]]
    stream_ids = [s["id"] for s in data["streams"]]
    tags = [u["tag"] for u in data["equipment"] if u["category"] != "logic"]

    # basis and standards
    findings = psr_basis.check_basis(basis)
    errs = [f for f in findings if f["level"] == "error"]
    warns = [f for f in findings if f["level"] == "warning"]
    checks.append(_check("BASIS-COMPLETE", "Engineering basis is complete (objective, audience, maturity, boundaries, cases, criteria, "
                         "sources, model, conventions, classification, assumptions, exclusions, data gaps)",
                         "FAIL" if errs else ("WARN" if warns else "PASS"),
                         "; ".join("%s %s" % (f["code"], f["message"]) for f in (errs + warns)[:12]) or "no findings"))
    sf = psr_basis.check_standards(standards)
    checks.append(_check("STD-REGISTER", "Standards register has edition, status, source URL, date checked, application, class, "
                         "compliance statement and no unsupported compliance claim", "FAIL" if sf else "PASS",
                         "; ".join(f["message"] for f in sf[:12]) or "%d standards registered" % len(standards)))
    n_open = psr_draft.count_unconfirmed(basis)
    is_draft = str(basis.get("basis_status", "")).lower() == "draft"
    readiness = str((basis.get("document") or {}).get("readiness", "")).upper()
    if is_draft:
        status = "FAIL" if readiness not in ("SCREENING",) else "WARN"
        detail = "draft basis: %d items marked TO BE CONFIRMED; readiness must stay SCREENING until a person confirms the basis" % n_open
    elif n_open:
        status, detail = "FAIL", "basis is not marked draft but %d items still read TO BE CONFIRMED" % n_open
    else:
        status, detail = "PASS", "basis confirmed (no open items)"
    checks.append(_check("BASIS-STATUS", "Engineering basis is confirmed, or the report is labelled as a draft screening", status, detail))
    ages = []
    for row in standards or []:
        try:
            ages.append((datetime.date.today() - datetime.date.fromisoformat(str(row.get("date_checked"))[:10])).days)
        except ValueError:
            ages.append(None)
    stale = [a for a in ages if a is None or a > STANDARDS_MAX_AGE_DAYS]
    default_reg = str(basis.get("standards_register_source", "")).startswith("toolkit default")
    checks.append(_check("STD-VERIFY", "Standards editions were verified recently (max %d days) and are not the unverified toolkit default" % STANDARDS_MAX_AGE_DAYS,
                         "WARN" if (stale or default_reg) else "PASS",
                         "%s%d of %d register rows older than %d days or without a date; re-verify editions before issue" % (
                             "default register shipped with the toolkit (verify before issue); " if default_reg else "", len(stale), len(ages),
                             STANDARDS_MAX_AGE_DAYS)
                         if (stale or default_reg) else "all rows checked within %d days" % STANDARDS_MAX_AGE_DAYS))
    std = data["meta"].get("standard_conditions", {})
    bstd = (basis.get("conventions") or {}).get("standard_conditions", {})
    same = std and bstd and abs(std.get("t_c", 1e9) - bstd.get("t_c", -1e9)) < 1e-9 and abs(std.get("p_bara", 1e9) - bstd.get("p_bara", -1e9)) < 1e-9
    checks.append(_check("BASIS-STDCOND-MATCH", "Standard conditions used in the extraction equal the basis convention",
                         "PASS" if same else "FAIL", "data %s vs basis %s" % (std, bstd)))

    # data integrity
    dup = len(stream_ids) - len(set(stream_ids))
    missing_results = [(s["id"], c) for s in data["streams"] for c in cases if not (s["results"].get(c) or {}).get("available")]
    missing_unit = [(u["tag"], c) for u in data["equipment"] for c in cases if c not in u["results"]]
    checks.append(_check("DATA-CASES", "Every stream and unit has results for every operating case", "FAIL" if (missing_results or missing_unit) else "PASS",
                         "missing stream results %s; missing unit results %s" % (missing_results[:5], missing_unit[:5]) if (missing_results or missing_unit) else
                         "%d streams x %d cases, %d units" % (len(stream_ids), len(cases), len(data["equipment"]))))
    known = set(u["tag"] for u in data["equipment"])
    bad_ref = [s["id"] for s in data["streams"] for t in (s["producers"] + s["consumers"]) if t not in known]
    multi = [s["id"] for s in data["streams"] if len(s["consumers"]) > 1]
    checks.append(_check("DATA-TOPOLOGY", "Stream ids unique; producers/consumers reference registered equipment", "FAIL" if (dup or bad_ref) else "PASS",
                         "duplicate ids %d; bad references %s" % (dup, bad_ref[:5])))
    checks.append(_check("DATA-MULTICONSUMER", "Streams consumed by more than one unit (area balances count them once per area)",
                         "WARN" if multi else "PASS", "streams: %s" % multi[:10] if multi else "none"))
    labels = [s["label"] for s in data["streams"]]
    checks.append(_check("DATA-LABELS", "Stream labels are unique", "FAIL" if len(labels) != len(set(labels)) else "PASS", ""))
    audit = psr_connectivity.audit_connectivity(data)
    wired = (data.get("meta") or {}).get("wiring_applied") or []
    names = [("%s (%s)" % (f["name"], f["code"])) for f in audit["findings"] if f["severity"] == psr_connectivity.SEVERITY_ERROR][:6]
    checks.append(_check("TOPO-CONNECTIVITY", "Every unit is connected with streams (recycles have a tear stream, columns a feed, no isolated "
                         "streams); drawings and stream tables match the wiring of the solved model", audit["status"],
                         ("%d error(s), %d warning(s) %s; first: %s; repair candidates in data/connectivity_audit.csv%s" % (
                             audit["summary"]["n_error"], audit["summary"]["n_warning"], audit["summary"]["by_code"], names,
                             "; %d wiring repair(s) applied" % len(wired) if wired else ""))
                         if audit["status"] != "PASS" else
                         "all %d units connected%s" % (len(data["equipment"]), "; %d wiring repair(s) applied" % len(wired) if wired else "")))

    generic = [u["tag"] for u in data["equipment"] if u["category"] == "other"]
    checks.append(_check("EQ-GENERIC-SYMBOL", "Every equipment type has a dedicated schematic symbol (unknown types are drawn as a flagged generic box)",
                         "WARN" if generic else "PASS", "%d units drawn with the generic symbol: %s" % (len(generic), generic[:10]) if generic else "none"))
    eos_list = (data.get("meta") or {}).get("eos_list") or []
    conflicts = (data.get("meta") or {}).get("component_conflicts") or []
    checks.append(_check("MODEL-EOS-COMPONENTS", "Single thermodynamic model and unambiguous component definitions across all areas",
                         "WARN" if (len(eos_list) > 1 or conflicts) else "PASS",
                         "; ".join(([("several EOS classes in one model: %s (the basis must state which area uses which)" % ", ".join(eos_list))] if len(eos_list) > 1 else []) + conflicts[:4]) or "one EOS, one definition per component"))
    # compositions
    worst, where = 0.0, ""
    pworst, pwhere = 0.0, ""
    undefined_x = []
    for s in data["streams"]:
        for c in cases:
            r = s["results"].get(c) or {}
            if not r.get("available") or r.get("zero_flow"):
                continue
            tot = sum((r.get("z") or {}).values())
            if abs(tot - 1.0) > worst:
                worst, where = abs(tot - 1.0), "%s/%s" % (s["id"], c)
            for p in r.get("phases", []):
                if (p.get("mole_fraction") or 0.0) > 1e-12:
                    if p.get("x") is None:
                        undefined_x.append("%s/%s/%s (%.1e kg/h)" % (s["id"], c, p["type"], r.get("mass_kg_h") or 0.0))
                        if (r.get("mass_kg_h") or 0.0) >= 1e-3:
                            pworst, pwhere = 1.0, "%s/%s/%s undefined" % (s["id"], c, p["type"])
                        continue
                    xs = sum(v for v in (p.get("x") or {}).values() if v is not None)
                    if abs(xs - 1.0) > pworst:
                        pworst, pwhere = abs(xs - 1.0), "%s/%s/%s" % (s["id"], c, p["type"])
    checks.append(_check("COMP-SUM", "Overall composition sums to 1 (tolerance %.0e)" % COMP_SUM_TOL, "PASS" if worst <= COMP_SUM_TOL else "FAIL",
                         "max |sum-1| = %.2e at %s" % (worst, where)))
    zero = [(s["id"], c) for s in data["streams"] for c in cases if (s["results"].get(c) or {}).get("zero_flow")]
    checks.append(_check("DATA-ZERO-FLOW", "Streams with an undefined composition (no moles in the solver fluid) are reported as n/a, not as zero",
                         "WARN" if zero else "PASS", "%d stream/case pairs: %s" % (len(zero), sorted({z[0] for z in zero})[:12]) if zero else "none"))
    raw_bad = [(s["id"], c) for s in data["streams"] for c in cases
               if (s["results"].get(c) or {}).get("phase_beta_sum_raw") is not None and abs(s["results"][c]["phase_beta_sum_raw"] - 1.0) > 1e-6]
    checks.append(_check("PHASE-FRACTION-RAW", "Raw solver phase fractions sum to 1 (otherwise the extractor normalised them)", "WARN" if raw_bad else "PASS",
                         "%d stream/case pairs normalised, e.g. %s" % (len(raw_bad), raw_bad[:5]) if raw_bad else "all closed"))
    checks.append(_check("PHASE-COMP-SUM", "Phase compositions sum to 1 for phases with non-zero fraction (tolerance %.0e)" % PHASE_SUM_TOL,
                         "PASS" if pworst <= PHASE_SUM_TOL else "FAIL", "max |sum-1| = %.2e at %s" % (pworst, pwhere)))
    checks.append(_check("PHASE-COMP-UNDEFINED", "Degenerate phase compositions (sum of x not 1 in a numerically empty phase) are shown as n/a",
                         "WARN" if undefined_x else "PASS", "%d phases, e.g. %s" % (len(undefined_x), undefined_x[:4]) if undefined_x else "none"))

    # balances
    tol = balances["tolerances"]
    bad_overall = [(c, k) for c in cases for k in ("mass_status", "energy_status") if balances["cases"][c]["overall"][k] == "FAIL"]
    bad_area = [(c, a["area"], k) for c in cases for a in balances["cases"][c]["areas"] for k in ("mass_status", "energy_status") if a[k] == "FAIL"]
    exceptions = {str(e["id"]): e for e in basis.get("known_balance_exceptions", [])}
    unit_fail = [f for f in balances["fail_items"] if f["level"] == "unit"]
    undocumented = [f for f in unit_fail if f["id"] not in exceptions]
    checks.append(_check("BAL-OVERALL", "Overall mass and energy balance within tolerance for every case (mass %.0e, energy %.0e); a documented 'overall' exception turns FAIL into WARN" % (tol["overall_mass_rel"], tol["energy_rel"]),
                         ("WARN" if "overall" in exceptions else "FAIL") if bad_overall else "PASS",
                         ("failed (documented in the basis): %s" if "overall" in exceptions else "failed: %s") % bad_overall if bad_overall else "all cases closed"))
    undocumented_area = [b for b in bad_area if b[1] not in exceptions]
    checks.append(_check("BAL-AREA", "Area balances within tolerance; documented exceptions are listed in the basis",
                         "FAIL" if undocumented_area else ("WARN" if bad_area else "PASS"),
                         "%d area checks outside tolerance, %d without documented explanation: %s" % (len(bad_area), len(undocumented_area), undocumented_area[:10])
                         if bad_area else "all areas closed"))
    checks.append(_check("BAL-UNIT", "Unit balances within tolerance; documented exceptions are listed in the basis",
                         "FAIL" if undocumented else ("WARN" if unit_fail else "PASS"),
                         "%d unit checks outside tolerance, %d without documented explanation: %s" %
                         (len(unit_fail), len(undocumented), sorted({f["id"] for f in undocumented})[:10])))
    note_items = balances.get("note_items", [])
    checks.append(_check("BAL-NOTES", "Residuals of NeqSim equipment whose energy or recycle bookkeeping is not a closure test "
                         "(separator/scrubber re-flash, component splitter, recycle tear) are listed as NOTE with their size",
                         "WARN" if note_items else "PASS",
                         "%d checks, e.g. %s" % (len(note_items), sorted({"%s %s" % (n["id"], n["check"]) for n in note_items})[:6])
                         if note_items else "none"))

    # diagrams
    if diagrams is not None:
        import psr_diagrams

        in_pfd_streams, in_pfd_tags = set(), set()
        for sheet in diagrams["pfd"]["sheets"]:
            inv = psr_diagrams.inventory_svg(sheet["svg"])
            in_pfd_streams |= set(inv["streams"])
            in_pfd_tags |= set(inv["tags"])
        miss_s = sorted(set(stream_ids) - in_pfd_streams)
        miss_t = sorted(set(tags) - in_pfd_tags)
        extra_s = sorted(in_pfd_streams - set(stream_ids))
        extra_t = sorted(in_pfd_tags - set(tags))
        checks.append(_check("PFD-STREAMS", "Every modelled stream appears on a PFD sheet with its number", "FAIL" if miss_s else "PASS",
                             "missing: %s" % miss_s[:10] if miss_s else "%d streams on %d sheets" % (len(stream_ids), len(diagrams["pfd"]["sheets"]))))
        checks.append(_check("PFD-TAGS", "Every equipment tag appears on a PFD sheet", "FAIL" if miss_t else "PASS",
                             "missing: %s" % miss_t[:10] if miss_t else "%d tags placed" % len(tags)))
        checks.append(_check("PFD-NO-EXTRAS", "PFD contains no stream or tag that is absent from the data tables", "FAIL" if (extra_s or extra_t) else "PASS",
                             "extra streams %s; extra tags %s" % (extra_s[:5], extra_t[:5])))
        warn = diagrams["pfd"].get("warnings", [])
        checks.append(_check("PFD-WARNINGS", "Diagram generator reported no layout warnings", "WARN" if warn else "PASS", "; ".join(warn[:5])))
        checks.append(_check("BFD-PRESENT", "Block flow diagram files exist", "PASS" if diagrams.get("bfd") and os.path.exists(diagrams["bfd"]["svg"]) else "FAIL", ""))

    # exports
    if exports_dir:
        need = ["process_data.json", "process_data.xlsx", "stream_properties.csv", "stream_components.csv", "stream_phase_compositions.csv", "equipment.csv", "components.csv"]
        absent = [n for n in need if not os.path.exists(os.path.join(exports_dir, n))]
        detail = "all exports present"
        status = "FAIL" if absent else "PASS"
        if not absent:
            with open(os.path.join(exports_dir, "stream_components.csv"), newline="", encoding="utf-8") as fh:
                n_rows = sum(1 for _ in csv.reader(fh)) - 1
            expected = len(stream_ids) * len(data["components"]) * len(cases)
            if n_rows != expected:
                status, detail = "FAIL", "stream_components.csv has %d rows, expected %d" % (n_rows, expected)
            else:
                detail = "stream_components.csv has %d rows = %d streams x %d components x %d cases" % (n_rows, len(stream_ids), len(data["components"]), len(cases))
        else:
            detail = "missing %s" % absent
        checks.append(_check("EXPORTS", "Machine-readable exports complete (component flow rates for every stream, component and case)", status, detail))

    # document checks
    if docx_path and os.path.exists(docx_path):
        text, tables, pics, heads = _docx_inventory(docx_path)
        content_text = text
        hits = psr_basis.placeholder_hits(content_text)
        checks.append(_check("DOC-PLACEHOLDERS", "No placeholder text remains in the report", "FAIL" if hits else "PASS", str(hits[:5])))
        bad = BAD_TOKENS.findall(content_text)
        checks.append(_check("DOC-NO-SILENT-FILL", "No nan/inf/None tokens: unavailable data is shown as n/a", "FAIL" if bad else "PASS", str(sorted(set(b.lower() for b in bad))[:5])))
        idset = set(stream_ids)

        def cover(prefix, need_ids):
            found = None
            for cap, hdr, rows in tables:
                if cap.startswith("Table") and (" - " + prefix + " ") in cap:
                    ids = {r[0] for r in rows}
                    found = (found or set()) | ids if prefix == "B.1" else (found or []) + [ids]
            return found

        b1 = cover("B.1", idset)
        ok1 = b1 is not None and idset <= b1
        checks.append(_check("DOC-STREAM-REGISTER", "Stream register table lists every stream", "PASS" if ok1 else "FAIL", "" if ok1 else "B.1 incomplete or missing"))
        for prefix, label in (("B.2", "properties"), ("B.4", "composition")):
            lst = cover(prefix, idset) or []
            okp = len(lst) == len(cases) and all(idset <= s for s in lst)
            checks.append(_check("DOC-STREAM-%s" % prefix.replace(".", ""), "Stream %s table present for every case and complete" % label, "PASS" if okp else "FAIL",
                                 "%d tables for %d cases" % (len(lst), len(cases))))
        reg_ok = False
        for cap, hdr, rows in tables:
            if "Equipment register" in cap:
                reg_ok = set(tags) <= {r[0] for r in rows}
        checks.append(_check("DOC-EQUIPMENT-REGISTER", "Equipment register lists every unit tag", "PASS" if reg_ok else "FAIL", ""))
        pc = [c["name"] for c in data["components"] if c.get("is_tbp") or c.get("is_plus")]
        in_doc = [n for n in pc if n in text]
        checks.append(_check("DOC-HEAVY-ENDS", "Every characterised pseudo-component is documented with its properties", "PASS" if len(in_doc) == len(pc) else "FAIL",
                             "%d of %d in report" % (len(in_doc), len(pc))))
        nopics = [p for p in pics if not p[0].strip() or not p[1].startswith("Figure")]
        checks.append(_check("DOC-FIGURES", "Every figure has alt text and a caption", "FAIL" if nopics else "PASS", "%d figures checked" % len(pics)))
        refs = [s["reference"] for s in standards]
        miss_ref = [r for r in refs if r not in text]
        checks.append(_check("DOC-STANDARDS", "Standards register appears in the report", "FAIL" if miss_ref else "PASS", str(miss_ref[:5])))
        ready = (basis.get("document") or {}).get("readiness", "")
        checks.append(_check("DOC-READINESS", "Readiness status is stated in the report", "PASS" if ready and ready.upper() in text.upper() else "FAIL", ready))
        # executive summary traceability
        lines = text.split("\n")
        starts = [i for i, l in enumerate(lines) if l.strip() == "1 Executive summary"]
        start = starts[-1] if starts else None
        end = next((i for i, l in enumerate(lines) if start is not None and i > start and l.strip().startswith("2 Scope")), None)
        if start is None or end is None or report_kit is None or not facts_path or not os.path.exists(facts_path):
            checks.append(_check("DOC-EXEC-TRACE", "Executive-summary numbers are traceable", "FAIL", "executive summary or facts file not found"))
        else:
            with open(facts_path, encoding="utf-8") as fh:
                facts = json.load(fh)
            known_numbers = set()
            _numbers_from(facts["derived_facts"], known_numbers)
            _numbers_from([r["value"] for r in facts["resolved_references"]], known_numbers)
            _numbers_from(basis, known_numbers)
            quoted = report_kit.numbers_in_text(SECTION_REF_RE.sub("", "\n".join(lines[start + 1:end])))
            unmatched = [q for q in quoted if not any(report_kit._same_number(q, k) for k in known_numbers)]
            checks.append(_check("DOC-EXEC-TRACE", "Every number in the executive summary traces to extracted data, balances or the basis", "FAIL" if unmatched else "PASS",
                                 "%d numbers checked; untraced: %s" % (len(quoted), unmatched[:8])))
            checks.append(_check("DOC-REFS-RESOLVED", "All [[...]] data references resolved", "FAIL" if facts["unresolved_references"] else "PASS",
                                 str(facts["unresolved_references"][:5])))
    else:
        checks.append(_check("DOC-BUILT", "Report document exists", "SKIP", "pre-build validation"))

    n_fail = sum(1 for c in checks if c["status"] == "FAIL")
    n_warn = sum(1 for c in checks if c["status"] == "WARN")
    n_pass = sum(1 for c in checks if c["status"] == "PASS")
    return {"status": "FAIL" if n_fail else ("PASS WITH WARNINGS" if n_warn else "PASS"), "n_checks": len(checks), "n_pass": n_pass,
            "n_fail": n_fail, "n_warn": n_warn, "checks": checks}


def write_validation(result, out_dir):
    """Write validation_report.json and validation_report.md; returns both paths."""
    os.makedirs(out_dir, exist_ok=True)
    jp = os.path.join(out_dir, "validation_report.json")
    with open(jp, "w", encoding="utf-8") as fh:
        json.dump(result, fh, indent=1)
    mp = os.path.join(out_dir, "validation_report.md")
    lines = ["# Validation report", "", "Overall: **%s** (%d checks: %d pass, %d warn, %d fail)" % (result["status"], result["n_checks"], result["n_pass"], result["n_warn"], result["n_fail"]), "",
             "| Check | Result | Description | Detail |", "|---|---|---|---|"]
    for c in result["checks"]:
        lines.append("| %s | %s | %s | %s |" % (c["id"], c["status"], c["description"], c["detail"].replace("|", "/")[:300]))
    with open(mp, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    return jp, mp
