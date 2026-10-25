"""Render a frozen PVT-study dataset without running a thermodynamic calculation.

The semantic report blocks are shared by DOCX, PDF and HTML.  Figures are saved
as editable SVG/vector PDF and PNG previews.  Missing evidence is printed as
unavailable; a renderer never upgrades a numerical check to physical validation.
"""

import hashlib
import html
import json
import math
import os
import re
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx.shared import Mm, Pt, RGBColor  # noqa: E402
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

import psr_report  # noqa: E402
from pvt_data import quantity  # noqa: E402

NAVY = "#1f3a5f"
BLUE = "#2f718a"
ORANGE = "#aa5a20"
UNAVAILABLE = "Unavailable - evidence not supplied"


def _text(value):
    """Human-readable values; full precision remains in the frozen dataset."""
    if value is None:
        return UNAVAILABLE
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        if not math.isfinite(value):
            return "Non-finite value - invalid"
        return format(value, ".10g")
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, allow_nan=False, sort_keys=True)
    return str(value)


def _records(value):
    """Normalize list- or dictionary-keyed registers without losing identifiers."""
    if value is None:
        return []
    if isinstance(value, list):
        return [r if isinstance(r, dict) else {"value": r} for r in value]
    if isinstance(value, dict):
        if any(k in value for k in ("id", "status", "components", "model_id", "kind")):
            return [value]
        return [dict(v, register_id=k) if isinstance(v, dict)
                else {"register_id": k, "value": v} for k, v in value.items()]
    return [{"value": value}]


def _flatten(value, prefix=""):
    """Yield every scalar in a nested evidence register with its source path."""
    if isinstance(value, dict):
        if not value:
            yield prefix, "{}"
        for key, child in value.items():
            path = prefix + "." + str(key) if prefix else str(key)
            if key in ("reportMarkdown", "legacyReportMarkdown"):
                yield path, "Legacy text retained in the machine-readable dataset; normalized results are tabulated here"
                continue
            yield from _flatten(child, path)
    elif isinstance(value, (list, tuple)):
        if not value:
            yield prefix, "[]"
        for i, child in enumerate(value):
            yield from _flatten(child, "%s[%d]" % (prefix, i))
    else:
        yield prefix, _text(value)


def _safe_stem(value):
    """Return a portable output stem rather than trusting a study identifier."""
    return re.sub(r"[^A-Za-z0-9_-]+", "_", str(value)).strip("_")[:90] or "PVT_study"


def _safe_tokens(dataset):
    """Block unresolved report tokens before creating output files."""
    for path, value in _flatten(dataset):
        if isinstance(value, str) and re.search(r"\[\[[^\]]*\]\]", value):
            raise ValueError("Unresolved report token at %s" % path)


def _study(dataset):
    """Accept the study metadata contract and older document-control records."""
    return dataset.get("study") or dataset.get("document") or {}


def _readiness(dataset):
    """Display computed acceptance status ahead of supplied document metadata."""
    return dataset.get("status", _study(dataset).get("readiness", "UNASSESSED"))


def _is_synthetic(dataset):
    """Recognize explicit synthetic provenance without guessing its absence."""
    study = _study(dataset)
    if study.get("synthetic") or study.get("data_class") == "synthetic":
        return True
    return any(s.get("synthetic") or s.get("data_class") == "synthetic"
               or s.get("input_class") == "synthetic"
               for s in _records(dataset.get("samples")))


def _save_figure(fig, directory, stem, caption, source):
    """Write quantitative figures in editable and print-ready formats."""
    paths = {}
    for ext in ("svg", "pdf", "png"):
        path = directory / (stem + "." + ext)
        fig.savefig(path, dpi=180, bbox_inches="tight")
        paths[ext] = str(path)
    plt.close(fig)
    return dict(paths, id=stem, caption=caption, source=source)


def _box(ax, x, y, width, height, lines, color=NAVY):
    """Draw a readable conceptual apparatus block in normalized coordinates."""
    from matplotlib.patches import FancyBboxPatch

    ax.add_patch(FancyBboxPatch((x, y), width, height,
                              boxstyle="round,pad=0.012,rounding_size=0.012",
                              linewidth=1.3, edgecolor=color, facecolor="#f5f8fb"))
    ax.text(x + width / 2, y + height / 2, "\n".join(lines), ha="center",
            va="center", fontsize=9, color="#172b40", wrap=True)


def _arrow(ax, start, end, label=None, color=NAVY):
    """Draw one directed material-transfer relationship."""
    ax.annotate("", xy=end, xytext=start,
                arrowprops={"arrowstyle": "-|>", "color": color, "lw": 1.4})
    if label:
        label_y = (max(start[1], end[1]) + 0.015 if abs(start[0] - end[0]) < 0.01
                   else (start[1] + end[1]) / 2 + 0.035)
        ax.text((start[0] + end[0]) / 2, label_y,
                label, ha="center", va="bottom", fontsize=8, color=color)


def _protocol_figures(dataset, directory):
    """Render actual registered protocol kinds, labelled as conceptual apparatus."""
    out = []
    samples = _records(dataset.get("samples"))
    sample_ids = [str(s.get("id", s.get("sample_id", "unidentified"))) for s in samples]
    fig, ax = plt.subplots(figsize=(8.0, 3.0))
    ax.set(xlim=(0, 1), ylim=(0, 1))
    ax.axis("off")
    _box(ax, 0.025, 0.35, 0.25, 0.35,
         ["Registered samples", ", ".join(sample_ids[:4]) or "not supplied"])
    _box(ax, 0.39, 0.35, 0.24, 0.35, ["QC and preparation", "analysis / recombination"])
    _box(ax, 0.75, 0.35, 0.23, 0.35, ["Frozen fluid model", "parameters / export"])
    _arrow(ax, (0.28, 0.525), (0.38, 0.525))
    _arrow(ax, (0.64, 0.525), (0.74, 0.525))
    ax.text(0.5, 0.14, "Conceptual data lineage; see the sample and model registers for every identifier",
            ha="center", fontsize=8, color="#555555")
    out.append(_save_figure(fig, directory, "sample_lineage",
                           "Sample to final model lineage", "samples, analyses, models"))

    kinds = {}
    for experiment in _records(dataset.get("experiments")):
        kind = str(experiment.get("kind", experiment.get("type", ""))).upper()
        canonical = {"CME": "CCE", "CONSTANTMASSEXPANSION": "CCE",
                     "CONSTANTMASS_EXPANSION": "CCE", "CONSTANTVOLUMEDEPLETION": "CVD",
                     "DIFFERENTIALLIBERATION": "DLE", "DL": "DLE",
                     "SEPARATORTEST": "SEPARATOR"}.get(kind.replace(" ", ""), kind)
        if canonical in ("CCE", "CVD", "DLE", "SEPARATOR"):
            kinds.setdefault(canonical, []).append(experiment)
    for kind, experiments in kinds.items():
        fig, ax = plt.subplots(figsize=(8.0, 3.2))
        ax.set(xlim=(0, 1), ylim=(0, 1))
        ax.axis("off")
        ids = ", ".join(str(e.get("id", e.get("experiment_id", "unidentified")))
                        for e in experiments)
        if kind == "CCE":
            _box(ax, 0.33, 0.27, 0.34, 0.45,
                 ["Closed equilibrium cell", "gas and liquid retained", "fixed total composition"])
            _arrow(ax, (0.12, 0.5), (0.31, 0.5), "reduce pressure")
            ax.text(0.5, 0.83, "CCE / CME", fontsize=13, fontweight="bold", ha="center")
        elif kind == "CVD":
            _box(ax, 0.3, 0.27, 0.34, 0.45,
                 ["Fixed volume cell", "equilibrate at next pressure", "retain liquid"])
            _box(ax, 0.76, 0.4, 0.2, 0.24, ["Removed gas", "inventory ledger"])
            _arrow(ax, (0.65, 0.52), (0.75, 0.52))
            ax.text(0.5, 0.83, "CVD", fontsize=13, fontweight="bold", ha="center")
        elif kind == "SEPARATOR" and experiments[0].get("stages"):
            stages = experiments[0]["stages"]
            cols = min(4, len(stages))
            row_count = int(math.ceil(len(stages) / cols))
            fig.set_size_inches(8.0, max(3.2, 1.6 + row_count * 1.6))
            width, height = 0.8 / cols, 0.3 / row_count
            for i, stage in enumerate(stages):
                column, row = i % cols, i // cols
                x = 0.025 + column * 0.95 / cols
                y = 0.68 - row * 0.6 / row_count - height
                pressure = quantity(stage.get("pressure"), "bara")
                temperature = quantity(stage.get("temperature"), "K")
                name = stage.get("label") or ("Final stage" if i == len(stages) - 1 else "Separator %d" % (i + 1))
                _box(ax, x, y, width, height, [name, "P = %s bara" % _text(pressure),
                                              "T = %s K" % _text(temperature)])
                _arrow(ax, (x + width / 2, y + height + 0.02),
                       (x + width / 2, y + height + 0.16 / row_count), "gas")
                if column < cols - 1 and i < len(stages) - 1:
                    _arrow(ax, (x + width + 0.01, y + height / 2),
                           (x + 0.95 / cols - 0.01, y + height / 2))
        else:
            names = (["Depletion step i", "Depletion step i + 1", "Final retained liquid"]
                     if kind == "DLE" else ["Separator stage i", "Next separator stage", "Stock tank"])
            for x, label in zip((0.03, 0.38, 0.73), names):
                _box(ax, x, 0.33, 0.23, 0.35, [label, "retain liquid"])
                _arrow(ax, (x + 0.115, 0.69), (x + 0.115, 0.9), "removed gas")
            _arrow(ax, (0.27, 0.5), (0.37, 0.5))
            _arrow(ax, (0.62, 0.5), (0.72, 0.5))
        ax.text(0.5, 0.14, "Registered experiments: " + ids, ha="center", fontsize=8)
        e = experiments[0]
        pressures = e.get("pressures_bara", e.get("stage_pressures_bara", []))
        if kind == "SEPARATOR" and e.get("stages"):
            pressures = [quantity(stage.get("pressure"), "bara") for stage in e["stages"]]
        condition = "T = %s K" % _text(e.get("temperature_k"))
        if pressures:
            condition += "; P = %s to %s bara; %d pressure steps" % (
                _text(pressures[0]), _text(pressures[-1]), len(pressures))
        ax.text(0.5, 0.035, "Conceptual protocol; " + condition + "; complete conditions in the register",
                ha="center", fontsize=7.6, color="#555555")
        out.append(_save_figure(fig, directory, "protocol_" + kind.lower(),
                               kind + " experiment protocol", "experiments"))
    return out


def _observations(dataset):
    """Return long-format observations from the canonical study contract."""
    experiments = {e.get("id"): e for e in _records(dataset.get("experiments"))}
    out = []
    observed = set()
    for observation in _records(dataset.get("observations")):
        experiment_id = observation.get("experiment_id")
        experiment = experiments.get(experiment_id, {})
        prop = observation.get("property", observation.get("observable"))
        values = observation.get("values")
        if not isinstance(values, list):
            out.append(observation)
            observed.add((experiment_id, prop))
            continue
        pressures = experiment.get("pressures_bara", [])
        predictions = experiment.get("predictions", {}).get("values", {})
        baseline = experiment.get("baseline_predictions", {}).get("values", {})
        aliases = {"relative_volume": "relativeVolume", "liquid_dropout_percent": "liquidRelativeVolume",
                   "gas_z_factor": "gasZ", "rs_sm3_sm3": "rsSm3PerSm3", "bo_m3_sm3": "boM3PerSm3",
                   "oil_density_kg_m3": "oilDensityKgPerM3"}
        property_key = aliases.get(prop, prop)
        if prop == "viscosity_pa_s" and observation.get("phase") in ("gas", "oil", "aqueous"):
            property_key = observation["phase"] + "ViscosityPaS"
        predicted = observation.get("predicted_values", predictions.get(property_key, []))
        original = observation.get("baseline_values", baseline.get(property_key, []))
        uncertainty = observation.get("uncertainty")
        for i, value in enumerate(values):
            point = dict(observation)
            point.pop("values", None)
            point["measured"] = value
            point["source_kind"] = observation.get("source_kind", experiment.get("source_kind"))
            point["pressure_bara"] = pressures[i] if i < len(pressures) else None
            point["predicted"] = predicted[i] if i < len(predicted) else None
            point["baseline"] = original[i] if i < len(original) else None
            if isinstance(uncertainty, list):
                point["uncertainty"] = uncertainty[i] if i < len(uncertainty) else None
            out.append(point)
        observed.add((experiment_id, property_key))
    units = {"relativeVolume": "m3/m3 at saturation", "liquidRelativeVolume": "% of saturation volume",
             "oilDensityKgPerM3": "kg/m3", "gasZ": "dimensionless", "mixtureZ": "dimensionless",
             "yFactor": "dimensionless", "cumulativeMolePercentDepleted": "mol% of initial inventory",
             "boM3PerSm3": "m3/Sm3", "bgM3PerSm3": "m3/Sm3", "rsSm3PerSm3": "Sm3/Sm3",
             "gasViscosityPaS": "Pa s", "oilViscosityPaS": "Pa s", "aqueousViscosityPaS": "Pa s"}
    for experiment in experiments.values():
        pressures = experiment.get("pressures_bara", [])
        for prop, values in experiment.get("predictions", {}).get("values", {}).items():
            if prop not in units or not isinstance(values, list) or (experiment.get("id"), prop) in observed:
                continue
            if not any(isinstance(value, (int, float)) and math.isfinite(value) for value in values):
                continue
            for i, value in enumerate(values):
                out.append({"experiment_id": experiment.get("id"), "property": prop,
                            "unit": units[prop], "predicted": value,
                            "pressure_bara": pressures[i] if i < len(pressures) else None,
                            "role": experiment.get("role"), "source_kind": experiment.get("source_kind")})
    return out


def _number(record, names):
    """Read a finite real value, preserving absent-phase and failed observations."""
    for name in names:
        value = record.get(name)
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return float(value) if math.isfinite(value) else None
    return None


def make_figures(dataset, out_dir):
    """Build figures from supplied records only, including failed-data annotations."""
    directory = Path(out_dir)
    directory.mkdir(parents=True, exist_ok=True)
    with plt.rc_context({"font.family": "DejaVu Sans", "font.size": 9,
                         "svg.fonttype": "none", "axes.spines.top": False,
                         "axes.spines.right": False}):
        figures = _protocol_figures(dataset, directory)
        parameters = dataset.get("parameters") or {}
        components = parameters.get("components", []) if isinstance(parameters, dict) else []
        for mode, inventory in (("composition", components),
                                ("heavy_distribution", [c for c in components if c.get("isTBPFraction") or c.get("isPlusFraction")])):
            for start in range(0, len(inventory), 15):
                subset = inventory[start:start + 15]
                fig, ax = plt.subplots(figsize=(7.5, max(3.2, 1.2 + 0.29 * len(subset))))
                names = [str(c.get("name", "unidentified")) for c in subset]
                fractions = [_number(c, ("moleFraction",)) for c in subset]
                ax.barh(list(range(len(subset))), [float("nan") if v is None else v for v in fractions], color=NAVY)
                ax.set_yticks(list(range(len(subset))), labels=names)
                ax.invert_yaxis()
                ax.set_xlabel("Mole fraction of the total final fluid (mol/mol)")
                ax.grid(axis="x", alpha=0.2)
                title = "Final model component distribution" if mode == "composition" else "Characterized pseudo-component distribution"
                ax.set_title(title, loc="left", fontsize=11)
                fig.tight_layout()
                figures.append(_save_figure(fig, directory, mode + "_%03d" % start,
                                             title + "; model values, not additional analytical measurements",
                                             "parameters.components"))
        grouped = {}
        for row in _observations(dataset):
            prop = row.get("property", row.get("observable", "unspecified"))
            unit = row.get("unit", "unit not supplied")
            experiment = row.get("experiment_id", "unidentified")
            grouped.setdefault((experiment, prop, unit), []).append(row)
        for i, ((experiment, prop, unit), rows) in enumerate(grouped.items()):
            usable = [r for r in rows if _number(r, ("pressure_bara", "p_bara", "pressure")) is not None]
            if not usable:
                continue
            fig, ax = plt.subplots(figsize=(7.3, 3.8))
            x = [_number(r, ("pressure_bara", "p_bara", "pressure")) for r in usable]
            for key, label, style, color in (("baseline", "Untuned model", "--", "#777777"),
                                            ("predicted", "Final model", "-", NAVY),
                                            ("reloaded", "E300 reloaded", ":", BLUE)):
                values = [_number(r, (key, key + "_value")) for r in usable]
                if any(v is not None for v in values):
                    ax.plot(x, [float("nan") if v is None else v for v in values],
                            style, color=color, label=label, linewidth=1.5)
            for role, marker in (("calibration", "o"), ("validation", "s"), ("holdout", "s")):
                measured = [(r, xx) for r, xx in zip(usable, x)
                            if r.get("role", "calibration") == role
                            and _number(r, ("measured", "measured_value", "observed")) is not None]
                if measured:
                    mx = [v for _, v in measured]
                    my = [_number(r, ("measured", "measured_value", "observed")) for r, _ in measured]
                    err = [_number(r, ("uncertainty", "uncertainty_absolute")) or 0 for r, _ in measured]
                    label = "Measured " + role
                    if _is_synthetic(dataset) or any(r.get("source_kind") in (
                            "synthetic", "numerical", "synthetic_reference", "numerical_reference") for r, _ in measured):
                        label = "Synthetic reference " + role
                    ax.errorbar(mx, my, yerr=err, fmt=marker, ms=4, capsize=2,
                                mfc="white", color=ORANGE, label=label)
            failed = sum(r.get("status") in ("FAIL", "failed", "nonconverged") for r in rows)
            if failed:
                ax.text(0.02, 0.98, "%d failed observations retained in register" % failed,
                        transform=ax.transAxes, va="top", fontsize=8, color="#9b2525")
            ax.set_xlabel("Absolute pressure (bara)")
            ax.set_ylabel(str(prop) + " (" + str(unit) + ")")
            ax.set_title(str(experiment), loc="left", fontsize=11)
            ax.grid(alpha=0.2)
            handles, labels = ax.get_legend_handles_labels()
            if handles:
                ax.legend(handles, labels, fontsize=8, loc="best")
            fig.tight_layout()
            figures.append(_save_figure(fig, directory, "prediction_%03d" % i,
                                       "%s: %s; calibration and holdout roles remain distinct" % (experiment, prop),
                                       "observations"))
    return figures


def _bip_blocks(snapshot):
    """Find complete ordered matrices and split columns into readable blocks."""
    if isinstance(snapshot.get("capture"), dict):
        yield from _bip_blocks(snapshot["capture"])
    candidates = ("bip_matrix", "bip", "kij", "binaryInteractionParameters",
                  "binary_interaction_parameters", "binaryInteractions", "binaryInteraction")
    for key in candidates:
        matrix = snapshot.get(key)
        if isinstance(matrix, list) and matrix and all(isinstance(r, list) for r in matrix):
            n = len(matrix)
            if not all(len(row) == n for row in matrix):
                raise ValueError("BIP matrix must be square")
            for start in range(0, n, 6):
                columns = list(range(start, min(n, start + 6)))
                header = ["Component i"] + ["j = %d" % j for j in columns]
                rows = [[str(i)] + [_text(matrix[i][j]) for j in columns] for i in range(n)]
                yield key, header, rows
    for phase in snapshot.get("phases", []):
        pairs = phase.get("binaryInteractionParameters", [])
        if not pairs or not all(isinstance(p, dict) and "i" in p and "j" in p for p in pairs):
            continue
        n = len(snapshot.get("components", []))
        index = {(pair["i"], pair["j"]): pair for pair in pairs}
        if len(index) != n * n or len(pairs) != n * n or set(index) != {(i, j) for i in range(n) for j in range(n)}:
            raise ValueError("BIP register must include every ordered pair and zero entries")
        for parameter in ("kij", "kijT1"):
            for start in range(0, n, 6):
                columns = list(range(start, min(n, start + 6)))
                header = ["Component i"] + ["j = %d" % j for j in columns]
                rows = [[str(i)] + [_text(index[(i, j)].get(parameter)) for j in columns]
                        for i in range(n)]
                yield "phase %s %s" % (phase.get("index", "unassigned"), parameter), header, rows


def _component_blocks(snapshot):
    """Transpose full component inventories into readable, indexed table blocks."""
    if isinstance(snapshot.get("capture"), dict):
        yield from _component_blocks(snapshot["capture"])
    components = snapshot.get("components", [])
    if not components:
        return
    keys = []
    for component in components:
        for key in component:
            if key not in keys:
                keys.append(key)
    for start in range(0, len(components), 5):
        subset = components[start:start + 5]
        header = ["Property and unit"] + ["Component %d" % i for i in range(start, start + len(subset))]
        rows = [[key] + [_text(component.get(key)) for component in subset] for key in keys]
        yield header, rows


def _scalar_parameters(value):
    """Avoid duplicating component inventories and matrices in scalar tables."""
    if isinstance(value, dict):
        out = {}
        for key, child in value.items():
            if key == "portableFluidJson":
                continue
            if key in ("components", "eosComponentParameters", "binaryInteractionParameters", "bip_matrix", "bip", "kij") and isinstance(child, list):
                out[key] = "Complete values in the indexed component or BIP appendix below"
            else:
                out[key] = _scalar_parameters(child)
        return out
    if isinstance(value, list):
        return [_scalar_parameters(child) for child in value]
    return value


def _table_block(title, value, path):
    """Represent a register as path/value pairs with full evidence coverage."""
    rows = list(_flatten(value, path))
    return {"type": "table", "title": title, "header": ["Dataset reference", "Value"],
            "rows": rows or [(path, UNAVAILABLE)]}


def _experiment_blocks(dataset):
    """Render normalized experiment grids rather than repeating array paths."""
    experiments = _records(dataset.get("experiments"))
    if not experiments:
        yield _table_block("Experiment definitions and results", None, "experiments")
    for index, experiment in enumerate(experiments):
        identifier = experiment.get("id", str(index))
        metadata = {key: value for key, value in experiment.items() if key != "predictions"}
        yield _table_block("Experiment %s definition" % identifier, metadata, "experiments[%d]" % index)
        predictions = experiment.get("predictions", {})
        values = predictions.get("values", {})
        summary = {key: value for key, value in predictions.items() if key != "values"}
        yield _table_block("Experiment %s result metadata" % identifier, summary,
                           "experiments[%d].predictions" % index)
        pressures = experiment.get("pressures_bara", predictions.get("pressuresBara", []))
        keys = [key for key, value in values.items() if isinstance(value, list)]
        for start in range(0, len(keys), 4):
            selected = keys[start:start + 4]
            n = max([len(pressures)] + [len(values[key]) for key in selected])
            rows = [[_text(pressures[i] if i < len(pressures) else None)] + [
                _text(values[key][i] if i < len(values[key]) else None) for key in selected]
                    for i in range(n)]
            yield {"type": "table", "title": "%s results in registered pressure order" % identifier,
                   "header": ["P absolute bara"] + selected, "rows": rows}


def _analysis_blocks(dataset):
    """Keep the laboratory composition basis beside its normalized model input."""
    analyses = _records(dataset.get("analyses"))
    if not analyses:
        yield _table_block("Analysis register", None, "analyses")
    known = {"name", "kind", "value", "raw_value", "mole_fraction", "molar_mass_kg_mol", "density_g_cm3"}
    for index, analysis in enumerate(analyses):
        identifier = analysis.get("id", str(index))
        metadata = {key: value for key, value in analysis.items() if key != "components"}
        yield _table_block("Analysis %s method provenance and basis" % identifier, metadata, "analyses[%d]" % index)
        rows = []
        extra = []
        for component in analysis.get("components", []):
            rows.append([_text(component.get("name")), _text(component.get("kind", "pure")),
                         _text(component.get("raw_value", component.get("value"))),
                         _text(component.get("mole_fraction")),
                         _text(component.get("molar_mass_kg_mol")), _text(component.get("density_g_cm3"))])
            remainder = {key: value for key, value in component.items() if key not in known}
            if remainder:
                extra.append(dict(remainder, name=component.get("name")))
        if rows:
            yield {"type": "table", "title": "%s raw basis %s and prepared mole fractions" % (
                identifier, analysis.get("basis", "unavailable")),
                   "header": ["Component", "Kind", "Raw value", "Prepared mol/mol", "Molar mass kg/mol", "Density g/cm3"],
                   "rows": rows}
        if extra:
            yield _table_block("Analysis %s additional component evidence" % identifier, extra,
                               "analyses[%d].components" % index)


def _observation_blocks(dataset, raw_only=False):
    """Keep observed, untuned, final and residual values together at each point."""
    experiments = {e.get("id"): e for e in _records(dataset.get("experiments"))}
    observations = _records(dataset.get("observations"))
    if not observations:
        yield _table_block("Observation and residual register", None, "observations")
    array_keys = {"values", "predicted_values", "baseline_values", "residuals", "uncertainty"}
    for index, observation in enumerate(observations):
        identifier = observation.get("id", str(index))
        metadata = {key: value for key, value in observation.items() if key not in array_keys}
        experiment = experiments.get(observation.get("experiment_id"), {})
        metadata["experiment_temperature_k"] = experiment.get("temperature_k")
        yield _table_block("Observation %s provenance and role" % identifier, metadata, "observations[%d]" % index)
        values = observation.get("values")
        if not isinstance(values, list):
            continue
        pressures = experiments.get(observation.get("experiment_id"), {}).get("pressures_bara", [])
        rows = []
        for i, value in enumerate(values):
            row = [_text(pressures[i] if i < len(pressures) else None), _text(value)]
            for key in (("uncertainty",) if raw_only else ("baseline_values", "predicted_values", "residuals", "uncertainty")):
                items = observation.get(key)
                number = items[i] if isinstance(items, list) and i < len(items) else items if key == "uncertainty" else None
                row.append(_text(number))
            rows.append(row)
        observed_label = "Synthetic ref" if observation.get("source_kind") == "synthetic_reference" else "Observed"
        title_kind = "raw input values" if raw_only else "values and absolute residuals"
        yield {"type": "table", "title": "%s %s %s in %s" % (
            identifier, observation.get("property", "observable"), title_kind, observation.get("unit", "unit unavailable")),
               "header": (["P bara", observed_label, "Uncertainty"] if raw_only else
                          ["P bara", observed_label, "Untuned", "Final", "Residual", "Uncertainty"]), "rows": rows}


def report_blocks(dataset, figures):
    """Create one traceable report representation for every output format."""
    study = _study(dataset)
    blocks = []

    def section(number, title, paragraph, entries):
        blocks.append({"type": "heading", "text": "%d %s" % (number, title)})
        if paragraph:
            blocks.append({"type": "paragraph", "text": paragraph})
        for label, key in entries:
            blocks.append(_table_block(label, dataset.get(key), key))

    readiness = _readiness(dataset)
    conclusion = study.get("conclusion") or (
        "The study readiness is %s. Review the recorded findings and export audit before using the model." % readiness)
    component_count = len((dataset.get("parameters") or {}).get("components", [])) if isinstance(dataset.get("parameters"), dict) else 0
    conclusion += " The frozen records contain %d samples, %d analyses, %d final components and %d experiments." % (
        len(_records(dataset.get("samples"))), len(_records(dataset.get("analyses"))), component_count,
        len(_records(dataset.get("experiments"))))
    if _is_synthetic(dataset):
        conclusion += " This is a synthetic workflow demonstration. Numerical replay does not establish experimental accuracy."
    section(1, "Executive summary", conclusion, [("Document and study control", "study")])
    section(2, "Scope objectives and evidence", study.get("objective", UNAVAILABLE),
            [("Sources and applicability", "sources"), ("Standards applicability register", "standards"),
             ("Assumptions", "assumptions"), ("Data gaps", "data_gaps")])
    section(3, "Sampling and representativeness",
            "Sample identities and field evidence define the model lineage. Unavailable metadata remains unavailable.",
            [("Sample register", "samples")])
    section(4, "Sample quality and acceptance",
            "Quality findings remain visible; a successful EOS calculation cannot establish sample representativeness.",
            [("Sample and analysis QC", "qc")])
    section(5, "Analytical methods and raw study data",
            "The following input tables precede model results. Raw PVT reference values and prepared compositions retain their source, basis, units, uncertainty and data role.", [])
    if dataset.get("synthetic_reference"):
        blocks.append(_table_block("Synthetic reference generation and frozen input identity",
                                   dataset["synthetic_reference"], "synthetic_reference"))
    blocks.extend(_analysis_blocks(dataset))
    blocks.extend(_observation_blocks(dataset, raw_only=True))
    for fig in figures:
        if fig["id"].startswith("composition_"):
            blocks.append({"type": "figure", "figure": fig})
    section(6, "Recombination and material balance",
            "Measured and adjusted recombination ratios require distinct provenance. Review closure before characterization.",
            [("Recombination contracts", "recombination"), ("Calculated recombination audit", "recombination_audit")])
    lineage = next((f for f in figures if f["id"] == "sample_lineage"), None)
    if lineage:
        blocks.append({"type": "figure", "figure": lineage})
    section(7, "PVT experiment program",
            "The registered protocol controls phase retention and removal. A TP scan is not a depletion experiment.", [])
    blocks.extend(_experiment_blocks(dataset))
    for fig in figures:
        if fig["id"].startswith("protocol_"):
            blocks.append({"type": "figure", "figure": fig})
    section(8, "Model selection and baseline",
            "EOS variants and mixing rules remain explicit. Parameters are not transferable between EOS variants without validation.", [])
    model_lineage = [{k: v for k, v in m.items() if k != "capture"} for m in _records(dataset.get("models"))]
    blocks.append(_table_block("Model lineage and snapshots", model_lineage, "models"))
    section(9, "Heavy fraction characterization",
            "Splitting and lumping must preserve the recorded mole, mass and heavy-fraction moments within declared tolerances.", [])
    characterization = dataset.get("characterization")
    if isinstance(characterization, dict):
        audit = {key: value for key, value in characterization.items() if key not in ("source_components", "final_components")}
        blocks.append(_table_block("Characterization audit", audit, "characterization"))
        for key in ("source_components", "final_components"):
            for header, rows in _component_blocks({"components": characterization.get(key, [])}):
                blocks.append({"type": "table", "title": "Characterization %s" % key.replace("_", " "),
                               "header": header, "rows": rows})
    else:
        blocks.append(_table_block("Characterization audit", characterization, "characterization"))
    for fig in figures:
        if fig["id"].startswith("heavy_distribution_"):
            blocks.append({"type": "figure", "figure": fig})
    section(10, "Calibration strategy and results",
            "Only calibration data may drive fitting. A finite objective alone does not establish convergence or identifiability.",
            [("Calibration record", "calibration")])
    section(11, "Independent validation and uncertainty",
            "Holdout records remain separate from fitting. Synthetic predictions are numerical fixtures rather than laboratory measurements.", [])
    blocks.extend(_observation_blocks(dataset))
    blocks.extend([_table_block(title, dataset.get(key), key) for title, key in (
        ("Validation evidence", "validation"), ("Metrics by observable and data role", "metrics"),
        ("Uncertainty and sensitivity", "uncertainty"))])
    for fig in figures:
        if fig["id"].startswith("prediction_"):
            blocks.append({"type": "figure", "figure": fig})
    section(12, "Thermodynamic and numerical consistency",
            "Individual Boolean checks and failure diagnostics are retained. Consistency is necessary but does not prove experimental accuracy.",
            [("Thermodynamic consistency checks", "consistency")])
    section(13, "Final model parameters and validity",
            "The machine-readable package retains full precision. The appendix lists every supplied effective parameter rather than fitted values alone.",
            [("Validity envelope", "validity")])
    section(14, "Export replay and compatibility",
            "E300 replay errors are distinct from laboratory-model errors. Target-simulator compatibility remains unverified unless an independent target-ingestion record is supplied.",
            [("Native and E300 export audit", "exports")])
    section(15, "Conclusions open gaps and review",
            conclusion, [("Findings and accountable review", "findings"), ("Review record", "review")])
    section(16, "Parameter appendices and reproducibility",
            "The dataset path in every table identifies its authoritative record. Binary interaction matrices preserve component ordering and zero entries.",
            [("Software and source revisions", "software"), ("Execution timings", "timings"),
             ("Rerun instructions", "rerun"), ("Scientific and implementation references", "references")])
    parameters = dataset.get("parameters")
    printed = _scalar_parameters(parameters)
    blocks.append(_table_block("Final parameter register", printed, "parameters"))
    snapshots = _records(parameters)
    if not snapshots:
        snapshots = [m for m in _records(dataset.get("models")) if m.get("stage") == "final"]
    for i, snapshot in enumerate(snapshots):
        for header, rows in _component_blocks(snapshot):
            blocks.append({"type": "table", "title": "Model %s full component properties" % snapshot.get("id", i),
                           "header": header, "rows": rows})
        for phase in snapshot.get("phases", []):
            for header, rows in _component_blocks({"components": phase.get("eosComponentParameters", [])}):
                blocks.append({"type": "table", "title": "Model %s phase %s complete EOS component parameters" % (
                    snapshot.get("id", i), phase.get("index", "unavailable")), "header": header, "rows": rows})
        for key, header, rows in _bip_blocks(snapshot):
            blocks.append({"type": "table", "title": "Model %s complete %s matrix" % (snapshot.get("id", i), key),
                           "header": header, "rows": rows})
    blocks.append(_table_block("Figure register", [{"id": f["id"], "caption": f["caption"],
                                                     "source": f["source"]} for f in figures], "figures"))
    return blocks


class _LiteralResolver:
    """Use the shared report helpers without introducing unresolved fact tokens."""

    def expand(self, value):
        return str(value)


class PvtReportBuilder(psr_report.ReportBuilder):
    """PVT report using the established process-report table and caption helpers."""

    def __init__(self):
        self._init_document()
        self.res = _LiteralResolver()
        self.opt = {}
        for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"):
            style = self.doc.styles[name]
            style.font.color.rgb = RGBColor(0, 0, 0)
            properties = style.element.find(qn("w:pPr"))
            if properties is not None:
                border = properties.find(qn("w:pBdr"))
                if border is not None:
                    properties.remove(border)
        self.doc.styles["Normal"].font.size = Pt(11)
        for section in self.doc.sections:
            section.page_width, section.page_height = Mm(210), Mm(297)
            section.left_margin = section.right_margin = Mm(19)
            section.top_margin = section.bottom_margin = Mm(18)

    def build(self, dataset, blocks, path):
        """Save the DOCX from the same blocks used by the PDF and HTML renderers."""
        study = _study(dataset)
        self.doc.core_properties.title = study.get("title", "PVT fluid characterization study")
        self.doc.core_properties.author = study.get("prepared_by", "Not assigned")
        self.doc.core_properties.subject = "PVT study and final fluid model handover"
        self.doc.add_paragraph(study.get("title", "PVT fluid characterization study"), style="Title")
        self.P("Document %s   Revision %s   Status %s" % (
            study.get("id", "Not assigned"), study.get("revision", "Not assigned"),
            _readiness(dataset)))
        self.P("This report documents sample evidence, fluid characterization, experiment results and model export fidelity. "
               "Use the recorded validity envelope and open findings to assess whether the model is suitable for the intended decision.")
        if _is_synthetic(dataset):
            self.P("SYNTHETIC WORKFLOW DEMONSTRATION - numerical verification only", bold=True)
        footer = self.doc.sections[0].footer.paragraphs[0]
        footer.add_run("%s | Rev %s | " % (study.get("id", "PVT study"), study.get("revision", "unassigned")))
        psr_report._add_field(footer, "PAGE")
        for block in blocks:
            if block["type"] == "heading":
                self.H(block["text"])
            elif block["type"] == "paragraph":
                self.P(block["text"])
            elif block["type"] == "figure":
                figure = block["figure"]
                self.picture(figure["png"], 165, figure["caption"], figure["caption"])
            elif block["type"] == "table":
                n = len(block["header"])
                widths = [73, 99] if n == 2 else [28] + [144 / (n - 1)] * (n - 1)
                self.table(block["header"], block["rows"], caption=block["title"],
                           widths=widths, font=8.5, align=["left"] * n)
        self.doc.save(path)

    def picture(self, path, width_mm, alt, caption):
        """Keep every diagram and its numbered caption on the same page."""
        super().picture(path, width_mm, alt, caption)
        self.doc.paragraphs[-2].paragraph_format.keep_with_next = True
        self.doc.paragraphs[-1].paragraph_format.keep_together = True

    def table(self, *args, **kwargs):
        """Apply consistent borders, padding and row shading to shared tables."""
        table = super().table(*args, **kwargs)
        trailing = self.doc.paragraphs[-1]
        trailing.paragraph_format.space_before = Pt(0)
        trailing.paragraph_format.space_after = Pt(1)
        trailing.paragraph_format.line_spacing = Pt(1)
        trailing.add_run().font.size = Pt(1)
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                properties = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for side in ("top", "left", "bottom", "right"):
                    edge = OxmlElement("w:" + side)
                    edge.set(qn("w:val"), "single")
                    edge.set(qn("w:sz"), "4")
                    edge.set(qn("w:color"), "D9D9D9")
                    borders.append(edge)
                properties.append(borders)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    margin = OxmlElement("w:" + side)
                    margin.set(qn("w:w"), "75")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                properties.append(margins)
                if row_index > 0 and row_index % 2 == 0:
                    psr_report._shade(cell, "F5F8FB")
        return table


def _write_html(dataset, blocks, path):
    """Write a portable HTML report with safe escaped data and relative figures."""
    def esc(value):
        return html.escape(str(value))

    content = ["<!doctype html><html lang='en'><meta charset='utf-8'>",
               "<title>" + esc(_study(dataset).get("title", "PVT fluid characterization study")) + "</title>",
               "<style>body{font:16px/1.5 system-ui,sans-serif;max-width:1100px;margin:36px auto;padding:0 24px;color:#172b40}"
               "h1,h2{color:#111}table{border-collapse:collapse;width:100%;margin:16px 0 28px;table-layout:fixed}"
               "td,th{border:1px solid #d9d9d9;padding:8px;text-align:left;overflow-wrap:anywhere}"
               "th{background:#1f3a5f;color:white}tr:nth-child(even){background:#f5f8fb}"
               "img{max-width:100%}figcaption,caption{text-align:left;font-weight:600}figure{margin:24px 0}"
               "@media print{body{font-size:10pt}h2{break-after:avoid}thead{display:table-header-group}}</style><body>",
               "<h1>" + esc(_study(dataset).get("title", "PVT fluid characterization study")) + "</h1>"]
    for block in blocks:
        if block["type"] == "heading":
            content.append("<h2>" + esc(block["text"]) + "</h2>")
        elif block["type"] == "paragraph":
            content.append("<p>" + esc(block["text"]) + "</p>")
        elif block["type"] == "figure":
            figure = block["figure"]
            src = os.path.relpath(figure["svg"], Path(path).parent)
            content.append("<figure><img src='%s' alt='%s'><figcaption>%s</figcaption></figure>" %
                           (esc(src), esc(figure["caption"]), esc(figure["caption"])))
        else:
            content.append("<table><caption>%s</caption><thead><tr>%s</tr></thead><tbody>" %
                           (esc(block["title"]), "".join("<th>" + esc(h) + "</th>" for h in block["header"])))
            for row in block["rows"]:
                content.append("<tr>" + "".join("<td>" + esc(v) + "</td>" for v in row) + "</tr>")
            content.append("</tbody></table>")
    Path(path).write_text("\n".join(content) + "</body></html>\n", encoding="utf-8")


def _write_pdf(dataset, blocks, path):
    """Render a deterministic PDF from report blocks when ReportLab is installed.

    This PDF is generated directly from the canonical dataset, rather than
    pretending that a Word field-update/conversion occurred.
    """
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Image, KeepTogether, LongTable, Paragraph, SimpleDocTemplate, Spacer
    except ImportError:
        return {"status": "UNAVAILABLE", "reason": "ReportLab is not installed; DOCX and HTML are available"}

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("PVTBody", fontName="Helvetica", fontSize=9.5, leading=13,
                              spaceAfter=6, wordWrap="CJK"))
    styles.add(ParagraphStyle("PVTCell", parent=styles["PVTBody"], fontSize=7.3,
                              leading=9.5, spaceAfter=0, alignment=TA_LEFT))
    styles.add(ParagraphStyle("PVTHeader", parent=styles["PVTCell"], textColor=colors.white,
                              fontName="Helvetica-Bold"))
    styles["Title"].textColor = colors.black
    styles["Heading1"].textColor = colors.black
    story = []
    study = _study(dataset)

    def p(value, style="PVTBody"):
        return Paragraph(html.escape(str(value)).replace("\n", "<br/>"), styles[style])

    story.append(p(study.get("title", "PVT fluid characterization study"), "Title"))
    story.append(p("Document %s | Revision %s | Readiness %s" % (
        study.get("id", "Not assigned"), study.get("revision", "Not assigned"),
        _readiness(dataset))))
    if _is_synthetic(dataset):
        story.append(p("SYNTHETIC WORKFLOW DEMONSTRATION - numerical verification only"))
    for block in blocks:
        if block["type"] == "heading":
            story.append(p(block["text"], "Heading1"))
        elif block["type"] == "paragraph":
            story.append(p(block["text"]))
        elif block["type"] == "figure":
            from PIL import Image as PilImage

            figure = block["figure"]
            with PilImage.open(figure["png"]) as image:
                width, height = image.size
            story.append(KeepTogether([Image(figure["png"], width=170 * mm,
                                             height=170 * mm * height / width), p(figure["caption"])]))
        else:
            story.append(p(block["title"], "Heading2"))
            header = [p(v, "PVTHeader") for v in block["header"]]
            rows = [[p(v, "PVTCell") for v in row] for row in block["rows"]]
            n = len(header)
            widths = [74 * mm, 96 * mm] if n == 2 else [28 * mm] + [142 * mm / (n - 1)] * (n - 1)
            table = LongTable([header] + rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
            table.setStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(NAVY)),
                            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f8fb")]),
                            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d9d9d9")),
                            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                            ("LEFTPADDING", (0, 0), (-1, -1), 5),
                            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                            ("TOPPADDING", (0, 0), (-1, -1), 4),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 4)])
            story.extend([table, Spacer(1, 5 * mm)])

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(19 * mm, 12 * mm, "%s | Rev %s" % (study.get("id", "PVT study"), study.get("revision", "unassigned")))
        canvas.drawRightString(191 * mm, 12 * mm, str(doc.page))
        canvas.restoreState()

    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=19 * mm, leftMargin=19 * mm,
                            topMargin=18 * mm, bottomMargin=20 * mm,
                            title=study.get("title", "PVT fluid characterization study"),
                            author=study.get("prepared_by", "Not assigned"))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return {"status": "GENERATED", "renderer": "ReportLab canonical-dataset rendering", "path": str(path)}


def build_report(dataset, out_dir):
    """Build DOCX/PDF/HTML and editable figures from one frozen study dataset.

    The function performs no solver calls and does not mutate ``dataset``.
    ``pdf_status`` distinguishes direct rendering from any Word conversion.
    Physical acceptance and accountable review are never inferred by rendering.
    """
    if not isinstance(dataset, dict):
        raise ValueError("PVT report dataset must be a dictionary")
    _safe_tokens(dataset)
    directory = Path(out_dir)
    directory.mkdir(parents=True, exist_ok=True)
    canonical = json.dumps(dataset, sort_keys=True, ensure_ascii=False, allow_nan=False)
    digest = hashlib.sha256(canonical.encode("utf-8")).hexdigest()
    stem = _safe_stem(_study(dataset).get("id", "PVT_study"))
    figures = make_figures(dataset, directory / "figures")
    blocks = report_blocks(dataset, figures)
    docx_path = directory / (stem + ".docx")
    html_path = directory / (stem + ".html")
    pdf_path = directory / (stem + ".pdf")
    PvtReportBuilder().build(dataset, blocks, docx_path)
    _write_html(dataset, blocks, html_path)
    pdf_status = _write_pdf(dataset, blocks, pdf_path)
    register = directory / "figure_register.json"
    register.write_text(json.dumps(figures, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    result = {"docx": str(docx_path), "html": str(html_path),
              "pdf": str(pdf_path) if pdf_status["status"] == "GENERATED" else None,
              "pdf_status": pdf_status, "figures": figures, "figure_register": str(register),
              "dataset_sha256": digest, "physical_acceptance": "NOT INFERRED",
              "visual_inspection": "REQUIRED - render and inspect every page before issue"}
    (directory / "report_manifest.json").write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    return result
