"""Template-preserving Word report finishing and local visual QA.

Works on a separate output document, never the source. Rendering is explicit:
Word on Windows, otherwise an installed LibreOffice. Preview images are evidence
for human/agent review, not an automatic claim that the layout is approved.
"""
import argparse
import hashlib
import importlib.util
import json
import math
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def _property(parent, name, **attributes):
    """Set one OOXML property without accumulating duplicate elements."""
    element = parent.find(qn(name))
    if element is None:
        element = OxmlElement(name)
        parent.append(element)
    for key, value in attributes.items():
        element.set(qn("w:" + key), str(value))
    return element


def _field(paragraph, instruction, initial="1"):
    """Append a live Word field with a readable pre-update value."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_element = OxmlElement("w:instrText")
    instruction_element.set(qn("xml:space"), "preserve")
    instruction_element.text = " " + instruction + " "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = initial
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction_element, separate, text, end):
        run._r.append(element)


def _instructions(element):
    """Return both complex and simple field instructions."""
    return (element.xpath(".//w:instrText/text()")
            + element.xpath(".//w:fldSimple/@w:instr"))


def _is_caption(element):
    """Recognize a numbered table caption, including localized labels."""
    return any(re.search(r"\bSEQ\s+Table\b", text, re.I)
               for text in _instructions(element))


def _section_for(document, element):
    """Find the section whose page geometry governs a body element."""
    sections = iter(document.sections)
    section = next(sections)
    for sibling in document.element.body:
        if sibling is element:
            return section
        if sibling.tag == qn("w:p") and sibling.xpath("./w:pPr/w:sectPr"):
            section = next(sections, section)
    return section


def _preceding_caption(table):
    """Find a caption immediately before a table, ignoring empty paragraphs."""
    sibling = table._tbl.getprevious()
    while sibling is not None and sibling.tag == qn("w:p"):
        if _is_caption(sibling):
            return True
        if "".join(sibling.xpath(".//w:t/text()")).strip():
            return False
        sibling = sibling.getprevious()
    return False


def _body_start(document):
    """Leave the cover and contents untouched; start at the first body heading."""
    for paragraph in document.paragraphs:
        if paragraph.style.name.lower().startswith("heading"):
            return paragraph._p
    raise ValueError("No body heading found; refusing to restyle an unrecognized document.")


def _add_captions(document, body_elements, table_label):
    """Caption uncaptioned body tables using their existing section heading."""
    caption_style = next(
        (p.style for p in document.paragraphs if _is_caption(p._p)), None)
    if caption_style is None:
        caption_style = document.styles["Caption"]
    heading = ""
    added = []
    number = 0
    for element in body_elements:
        if element.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            paragraph = Paragraph(element, document)
            if paragraph.style.name.lower().startswith("heading"):
                heading = paragraph.text
        elif element.tag == qn("w:tbl"):
            from docx.table import Table
            table = Table(element, document)
            number += 1
            if _preceding_caption(table):
                continue
            if not heading:
                raise ValueError("Cannot infer a caption for a table without a section heading.")
            paragraph = document.add_paragraph(style=caption_style)
            paragraph.add_run(table_label + " ")
            _field(paragraph, r"SEQ Table \* ARABIC", str(number))
            paragraph.add_run(": " + heading)
            element.addprevious(paragraph._p)
            added.append(heading)
    return added


def _format_table(table, width):
    """Use readable cells, quiet rules and repeatable multipage headers."""
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    table.autofit = False
    properties = table._tbl.tblPr
    _property(properties, "w:tblW", w=round(width / 635), type="dxa")
    margins = _property(properties, "w:tblCellMar")
    for side, value in (("top", 75), ("bottom", 75), ("left", 95), ("right", 95)):
        _property(margins, "w:" + side, w=value, type="dxa")
    borders = _property(properties, "w:tblBorders")
    for side in ("top", "bottom", "insideH"):
        _property(borders, "w:" + side, val="single", sz=4, color="D5D9DD")
    for side in ("left", "right", "insideV"):
        _property(borders, "w:" + side, val="nil")
    columns = list(table.columns)
    old_widths = [column.width or Inches(1) for column in columns]
    if len(columns) == 2:
        lengths = [max((len(row.cells[index].text) for row in table.rows), default=0)
                   for index in range(2)]
        if lengths[1] > 3 * max(lengths[0], 1):
            old_widths = [3, 7]
        elif lengths[0] > 3 * max(lengths[1], 1):
            old_widths = [7, 3]
    total = sum(old_widths)
    widths = [round(width * value / total) for value in old_widths]
    widths[-1] += width - sum(widths)
    for column, value in zip(columns, widths):
        column.width = value
    compact = (len(table.rows) <= 8 and
               sum(len(cell.text) for row in table.rows for cell in row.cells) <= 2200)
    seen = set()
    for index, row in enumerate(table.rows):
        props = row._tr.get_or_add_trPr()
        if index == 0:
            _property(props, "w:tblHeader", val=1)
        # A truly oversized row must be able to continue rather than overflow.
        maximum = max((len(cell.text) for cell in row.cells), default=0)
        split = props.find(qn("w:cantSplit"))
        if maximum < 1800:
            _property(props, "w:cantSplit", val=1)
        elif split is not None:
            props.remove(split)
        for column_index, cell in enumerate(row.cells):
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            # gridSpan must retain the sum of its covered columns.
            span = cell._tc.grid_span
            cell.width = sum(widths[column_index:column_index + span])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index > 0 and index % 2 == 0:
                shading = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
                if shading is None:
                    _property(cell._tc.get_or_add_tcPr(), "w:shd",
                              val="clear", fill="F3F5F7")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.08
                paragraph.paragraph_format.keep_with_next = (
                    index == 0 or (compact and index < len(table.rows) - 1))
                paragraph.paragraph_format.widow_control = True
                for run in paragraph.runs:
                    run.font.size = Pt(10.5)
                    if index == 0:
                        run.bold = True
                if index > 0 and re.fullmatch(
                        r"\s*[-+~<>]?\s*[\d., %\u00a0\u2212]+\s*", paragraph.text):
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _add_page_numbers(document):
    """Keep existing footer content and numbering; add numbers only if absent."""
    seen = set()
    for section in document.sections:
        footer = section.footer
        if footer.part.partname in seen:
            continue
        seen.add(footer.part.partname)
        if any(re.search(r"\bPAGE\b", text, re.I)
               for text in _instructions(footer._element)):
            continue
        paragraph = footer.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _field(paragraph, "PAGE")
        paragraph.add_run(" / ")
        _field(paragraph, "NUMPAGES")
        for run in paragraph.runs:
            run.font.size = Pt(9)


def finish_document(source, output, table_label="Table", overwrite=False):
    """Produce a separately named, template-preserving formatting revision."""
    source, output = Path(source).resolve(), Path(output).resolve()
    if source == output or (output.exists() and source.samefile(output)):
        raise ValueError("Source and output must be different files.")
    if output.exists() and not overwrite:
        raise FileExistsError("Output already exists; use --overwrite explicitly: " + str(output))
    if output.suffix.lower() != ".docx":
        raise ValueError("Output must have a .docx extension.")
    if not table_label.strip():
        raise ValueError("Table label must not be empty.")
    document = Document(source)
    start = _body_start(document)
    elements = list(document.element.body)
    body = elements[elements.index(start):]
    original_text = [p.text for p in document.paragraphs]
    original_cells = [[cell.text for row in table.rows for cell in row.cells]
                      for table in document.tables]
    for section in document.sections:
        section.top_margin = max(section.top_margin, Inches(0.95))
        section.bottom_margin = max(section.bottom_margin, Inches(0.8))
    added = _add_captions(document, body, table_label)
    body_set = set(body)
    for paragraph in document.paragraphs:
        if paragraph._p not in body_set:
            continue
        name = paragraph.style.name.lower()
        if name.startswith("heading"):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
        elif name == "normal" and not paragraph._p.xpath(".//w:drawing"):
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.line_spacing = 1.15
            paragraph.paragraph_format.widow_control = True
            for run in paragraph.runs:
                if run.font.size in (Pt(11), Pt(12)):
                    run.font.size = None
            previous, following = paragraph._p.getprevious(), paragraph._p.getnext()
            if len(paragraph.text) < 120 and previous is not None and following is not None:
                from docx.text.paragraph import Paragraph
                if (previous.tag == qn("w:p") and following.tag == qn("w:p")
                        and Paragraph(previous, document).style.name.lower().startswith("heading")
                        and Paragraph(following, document).style.name.lower().startswith("heading")):
                    paragraph.paragraph_format.keep_with_next = True
    tables = [table for table in document.tables if table._tbl in body_set]
    warnings = []
    for table in tables:
        section = _section_for(document, table._tbl)
        width = section.page_width - section.left_margin - section.right_margin
        _format_table(table, width)
        # Bind the short table lead-in and caption to the header, not an entire table.
        sibling = table._tbl.getprevious()
        for _ in range(3):
            if sibling is None or sibling.tag != qn("w:p"):
                break
            from docx.text.paragraph import Paragraph
            paragraph = Paragraph(sibling, document)
            if len(paragraph.text) > 250:
                break
            paragraph.paragraph_format.keep_with_next = True
            sibling = sibling.getprevious()
        for row in table.rows[1:]:
            for cell in row.cells:
                if len(cell.text) > 600:
                    warnings.append("Long table cell requires editorial review; formatting "
                                    "does not summarize or remove technical evidence.")
    for paragraph in document.paragraphs:
        if any(re.search(r"\bSEQ\s+(Table|Figure|Equation)\b", text, re.I)
               for text in _instructions(paragraph._p)):
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.keep_with_next = _is_caption(paragraph._p)
    _add_page_numbers(document)
    # Formatting must not rewrite a study's numbers, wording, or evidence cells.
    remaining = iter(p.text for p in document.paragraphs)
    if not all(any(text == candidate for candidate in remaining) for text in original_text):
        raise RuntimeError("Formatting unexpectedly changed an original paragraph.")
    if original_cells != [[cell.text for row in table.rows for cell in row.cells]
                          for table in document.tables]:
        raise RuntimeError("Formatting unexpectedly changed table content.")
    _property(document.settings.element, "w:updateFields", val="true")
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return {
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "original_content_preserved": True,
        "body_tables": len(tables),
        "added_table_captions": added,
        "warnings": sorted(set(warnings)),
        "visual_review": "PENDING",
    }


def render_pdf(document, pdf):
    """Render a derivative locally, updating its fields but never the source."""
    document, pdf = Path(document).resolve(), Path(pdf).resolve()
    if sys.platform == "win32" and importlib.util.find_spec("win32com"):
        import win32com.client
        import pywintypes
        try:
            word = win32com.client.DispatchEx("Word.Application")
        except pywintypes.com_error as error:
            raise RuntimeError("Microsoft Word automation is unavailable: " + str(error)) from error
        word.Visible = False
        word.DisplayAlerts = 0
        word.AutomationSecurity = 3
        try:
            word.Options.UpdateLinksAtOpen = False
            doc = word.Documents.Open(str(document), ConfirmConversions=False,
                                      ReadOnly=False, AddToRecentFiles=False)
            try:
                doc.Repaginate()
                doc.Fields.Update()
                for toc in doc.TablesOfContents:
                    toc.Update()
                doc.Repaginate()
                # Refresh page references after the contents has changed size.
                for toc in doc.TablesOfContents:
                    toc.UpdatePageNumbers()
                doc.Save()
                doc.ExportAsFixedFormat(str(pdf), 17)
            finally:
                doc.Close(False)
        except pywintypes.com_error as error:
            raise RuntimeError("Microsoft Word rendering failed: " + str(error)) from error
        finally:
            word.Quit()
        backend = "Microsoft Word"
    else:
        executable = shutil.which("soffice") or shutil.which("libreoffice")
        if not executable:
            raise RuntimeError("PDF rendering requires Microsoft Word with pywin32 "
                               "or an installed LibreOffice; no PDF was rendered.")
        with tempfile.TemporaryDirectory(prefix="neqsim-report-render-") as temporary:
            profile = (Path(temporary) / "profile").as_uri()
            result = subprocess.run(
                [executable, "-env:UserInstallation=" + profile, "--headless",
                 "--convert-to", "pdf", "--outdir", temporary, str(document)],
                capture_output=True, text=True, timeout=120)
            produced = Path(temporary) / (document.stem + ".pdf")
            if result.returncode or not produced.exists():
                raise RuntimeError("LibreOffice PDF conversion failed: "
                                   + result.stdout + result.stderr)
            shutil.copyfile(produced, pdf)
        backend = "LibreOffice"
    if not pdf.exists() or pdf.read_bytes()[:4] != b"%PDF":
        raise RuntimeError("Renderer did not produce a valid PDF.")
    return backend


def preview_pdf(pdf, directory):
    """Render every page and a contact sheet; report basic geometry findings."""
    import fitz
    from PIL import Image, ImageDraw
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    previews, findings, thumbnails = [], [], []
    with fitz.open(pdf) as document:
        for number, page in enumerate(document, 1):
            output = directory / "page_{:03d}.png".format(number)
            page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(output)
            previews.append(str(output))
            words = page.get_text("words")
            if not words and not page.get_images():
                findings.append({"page": number, "issue": "Possibly blank page"})
            if any(word[0] < -1 or word[1] < -1 or word[2] > page.rect.width + 1
                   or word[3] > page.rect.height + 1 for word in words):
                findings.append({"page": number, "issue": "Text outside page bounds"})
            with Image.open(output) as image:
                thumbnail = image.convert("RGB")
            thumbnail.thumbnail((390, 552))
            tile = Image.new("RGB", (410, 580), "#e5e7eb")
            tile.paste(thumbnail, ((410 - thumbnail.width) // 2, 22))
            ImageDraw.Draw(tile).text((10, 5), "Page " + str(number), fill="black")
            thumbnails.append(tile)
        if not thumbnails:
            raise ValueError("Rendered PDF has no pages.")
        columns = min(4, len(thumbnails))
        contact = Image.new("RGB", (410 * columns, 580 * math.ceil(
            len(thumbnails) / columns)), "white")
        for index, thumbnail in enumerate(thumbnails):
            contact.paste(thumbnail, ((index % columns) * 410, (index // columns) * 580))
        sheet = directory / "contact_sheet.png"
        contact.save(sheet)
    for previous in directory.glob("page_*.png"):
        if re.fullmatch(r"page_\d{3,}\.png", previous.name) and str(previous) not in previews:
            previous.unlink()
    return {"pages": len(previews), "page_images": previews, "contact_sheet": str(sheet),
            "geometry_findings": findings, "visual_review": "PENDING"}


def main(argv=None):
    """CLI entry point: finish a separate DOCX and optionally render QA evidence."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="Existing generated Word report")
    parser.add_argument("--output", required=True, type=Path, help="Separate output .docx")
    parser.add_argument("--render", action="store_true", help="Export PDF and all-page previews")
    parser.add_argument("--table-label", default="Table", help="Caption label, e.g. Table or Tabell")
    parser.add_argument("--overwrite", action="store_true", help="Replace the derivative, never source")
    args = parser.parse_args(argv)
    qa = {}
    try:
        if not args.overwrite:
            companions = [args.output.with_suffix(".qa.json")]
            if args.render:
                companions.extend([args.output.with_suffix(".pdf"),
                                   args.output.parent / (args.output.stem + "_preview")])
            for path in companions:
                if path.exists():
                    raise FileExistsError("Output companion already exists; use --overwrite "
                                          "explicitly: " + str(path))
        qa = finish_document(args.source, args.output, args.table_label, args.overwrite)
        if args.render:
            pdf = args.output.with_suffix(".pdf")
            qa["renderer"] = render_pdf(args.output, pdf)
            qa.update(preview_pdf(pdf, args.output.parent / (args.output.stem + "_preview")))
        qa["output_sha256"] = hashlib.sha256(args.output.read_bytes()).hexdigest()
    except (OSError, ValueError, RuntimeError, ImportError) as error:
        print("ERROR: " + str(error), file=sys.stderr)
        if qa:
            qa["render_error"] = str(error)
            args.output.with_suffix(".qa.json").write_text(
                json.dumps(qa, indent=2), encoding="utf-8")
        return 2
    args.output.with_suffix(".qa.json").write_text(json.dumps(qa, indent=2), encoding="utf-8")
    print("Formatting revision: " + str(args.output))
    print("Visual review: PENDING - inspect the rendered page images before delivery.")
    for warning in qa["warnings"]:
        print("WARNING: " + warning)
    return 0


if __name__ == "__main__":
    sys.exit(main())
