"""Fail-closed validation of the SURF package: inputs, model, solver evidence, design checks, drawings and document.

Each check returns ``{id, description, status, detail}`` with status PASS, FAIL, WARN or INFO. A FAIL makes the package a
``NOT_READY`` diagnostic; open evidence (synthetic inputs, hand-offs, missing benchmarks) caps the label at ``SCREENING``. The software
never assigns ``REVIEWED``: that label needs a person and a recorded review (outside this tool).
"""
import json
import os
import re

BAD_TOKENS = re.compile(r"(?<![A-Za-z0-9_])(nan|NaN|-?inf|None)(?![A-Za-z0-9_])|\[\[|\]\]")
CONFORMITY = re.compile(r"complied|conforms|conformity (achieved|demonstrated)|compliant with", re.I)
REQUIRED_HEADINGS = ["1 Executive summary", "2 Scope", "3 Design basis", "4 Field", "5 Alternatives", "6 Generated SURF", "7 Fluids", "8 Flowline", "9 Riser", "10 Umbilical",
                     "11 Subsea equipment", "12 Materials", "13 Safety", "14 Installation", "15 Operations", "16 Interface", "17 MTO", "18 Sensitivity", "19 Conclusions", "20 Appendices"]


def _check(cid, description, status, detail=""):
    return {"id": cid, "description": description, "status": status, "detail": detail}


def check_data(ds):
    """Checks on the canonical dataset (inputs, model identity, solver evidence, design assessment)."""
    out = []
    errs = [f for f in ds["findings"] if f["level"] == "error"]
    warns = [f for f in ds["findings"] if f["level"] == "warning"]
    out.append(_check("V-INPUT", "Specification has no input errors", "FAIL" if errs else "PASS", "; ".join(f["message"] for f in errs[:3]) or ("%d warning(s)" % len(warns) if warns else "")))
    out.append(_check("V-CONCEPTS", "At least two meaningful concepts are compared", "PASS" if len(ds["concepts"]) >= 2 else "FAIL", ", ".join(ds["concepts"])))
    unconverged = []
    for cid, cases in ds["results"].items():
        for case, rec in cases.items():
            if not rec["derived"]["converged"]:
                unconverged.append((cid, case, "; ".join(rec["derived"].get("problems", []))[:90]))
    relevant = {ds["reference_concept"], ds["selected_concept"]}
    bad_rel = [u for u in unconverged if u[0] in relevant]
    out.append(_check("V-CONVERGE", "Every case of the reference/selected concept solved: coupling converged, process model converged, no retained hydraulic fallback",
                      "FAIL" if bad_rel else ("WARN" if unconverged else "PASS"),
                      "; ".join("%s/%s: %s" % u for u in unconverged[:4]) + (" (concepts that are not reference or selected are reported infeasible and excluded from ranking)" if unconverged and not bad_rel else "")))
    bad_status = []
    for cid, cases in ds["results"].items():
        for case, rec in cases.items():
            for dom in ("production", "injection"):
                for eid, e in rec[dom]["edges"].items():
                    st = (e.get("model_status") or "")
                    if re.search(r"fail|fallback|stale|not_converged|unconverged", st, re.I):
                        bad_status.append((cid, "%s/%s/%s: %s" % (cid, case, eid, st[:50])))
    out.append(_check("V-FALLBACK", "No retained hydraulic fallback or failed-edge status in the reference/selected concept", "FAIL" if any(b[0] in relevant for b in bad_status) else ("WARN" if bad_status else "PASS"),
                      "; ".join(b[1] for b in bad_status[:3])))
    relevant = {ds["reference_concept"], ds["selected_concept"]}
    hand_bad = [("%s %s" % (cid, h["id"])) for cid, hs in ds["hand_checks"].items() if cid in relevant for h in hs if h["status"] != "PASS"]
    out.append(_check("V-HAND", "Hydrostatic, friction, thermal and mass-closure hand checks agree with the solved values", "FAIL" if hand_bad else "PASS",
                      "%d checks; failed: %s" % (sum(len(v) for v in ds["hand_checks"].values()), ", ".join(hand_bad) or "none")))
    ms = ds.get("mesh_sensitivity")
    out.append(_check("V-MESH", "Spatial-sensitivity gate (study mesh against a finer mesh)", "WARN" if ms is None else ("PASS" if ms["status"] == "PASS" else "FAIL"),
                      "not run" if ms is None else "max dp %.3f bar, max dT %.2f K (tolerances %s)" % (ms["max_dp_bar"], ms["max_dt_k"], ms["tolerances"])))
    rk = [k for k, v in ds["cooldown_checks"].items() if v["rel_diff"] > 1.0e-3]
    out.append(_check("V-COOLDOWN", "Lumped cooldown time agrees with the independent RK4 integration", "FAIL" if rk else "PASS", ", ".join(rk)))
    fw = [c for c in ds.get("fluid_checks", []) if c["status"] != "PASS"]
    out.append(_check("V-FLUID", "SRK pure-water properties against tabulated water", "WARN" if fw else "PASS",
                      "; ".join("%s deviates %.1f%%" % (c["description"].split(" against")[0], 100 * c["deviation"]) for c in fw) or ""))
    dim = [("%s %s" % (cid, d["edge"])) for cid, ds_ in ds["diameter_checks"].items() for d in ds_ if d["status"] != "PASS"]
    out.append(_check("V-DIMENSIONS", "OD, wall and inside-diameter relations and corrosion allowance consistent", "FAIL" if dim else "PASS", ", ".join(dim)))
    rec_bad = [("%s %s" % (cid, r["item"])) for cid, regs in ds["registers"].items() for r in regs["reconciliation"] if r["status"] != "PASS"]
    out.append(_check("V-RECONCILE", "Model, register, design basis and cost basis agree (counts, lengths, diameters)", "FAIL" if rec_bad else "PASS", ", ".join(rec_bad[:4])))
    cost_missing = [cid for cid in ds["concepts"] if not ds["cost"][cid].get("total_usd")]
    out.append(_check("V-COST", "Cost basis present for every concept (currency, price year, class, scope)", "FAIL" if cost_missing else "WARN",
                      ", ".join(cost_missing) or "screening-level estimate from the existing SURFCostEstimator; not a maturity claim"))
    sel = ds["selected_concept"]
    rows = {r["concept"]: r for r in ds["comparison"]["rows"]}
    ranked_bad = sel is not None and rows[sel]["feasibility"] != "feasible"
    out.append(_check("V-FEASIBILITY", "Cost ranking only among concepts that are feasible in every case", "FAIL" if ranked_bad else "PASS", "selected: %s" % (sel or "none")))
    blocked = [k for cid, v in ds["structural"].items() for k, x in v.items() if x["status"] == "blocked"]
    open_s = [k for k, x in ds["structural"][ds["reference_concept"]].items() if not x["status"].startswith("screened")]
    out.append(_check("V-STRUCTURAL", "Structural kernel results: screened, blocked and not-assessed lines are explicit", "FAIL" if blocked else "WARN",
                      "%d line(s) screened with the DNV-ST-F101 kernel (synthetic inputs); %d not assessed (evidence missing); blocked: %s" % (
                          len(ds["structural"][ds["reference_concept"]]) - len(open_s), len(open_s), ", ".join(blocked) or "none")))
    conf = [s["id"] for s in ds["standards"] if CONFORMITY.search(str(s.get("status", "")))]
    out.append(_check("V-STANDARDS", "Standards register makes no conformity claim; text availability is stated", "FAIL" if conf else "WARN",
                      ", ".join(conf) or "%d standards listed; none read at clause level" % len(ds["standards"])))
    na = [c["id"] for cid in ds["criteria"] for c in ds["criteria"][cid] if c["overall"] == "NOT ASSESSED"]
    out.append(_check("V-CRITERIA", "Every acceptance criterion was assessed", "WARN" if na else "PASS", ", ".join(sorted(set(na)))))
    out.append(_check("V-BENCHMARK", "Independent multiphase benchmark against public/lab/field data", "WARN",
                      "none available in this toolkit: Beggs and Brill carries no general experimental qualification; hand checks cover single-phase hydrostatics, friction, thermal and mass only"))
    out.append(_check("V-ENERGY", "Independent energy/enthalpy closure of the network", "WARN", "conservative mixing is owned by the solver; only mass closure is independently re-computed"))
    out.append(_check("V-HANDOFFS", "Specialist hand-offs for unsupported analyses are listed with prepared inputs", "WARN" if ds["handoffs"] else "FAIL",
                      "%d open hand-off(s): %s" % (len(ds["handoffs"]), ", ".join(h["id"] for h in ds["handoffs"]))))
    out.append(_check("V-GAPS", "Data gaps and assumptions are recorded", "WARN" if ds["data_gaps"] else "INFO", "%d data gap(s), %d assumption(s)" % (len(ds["data_gaps"]), len(ds["assumptions"]))))
    return out


def check_drawings(ds, manifest):
    """Checks on the drawing package: layout QA, numbering, asset identity agreement and title-block responsibilities."""
    out = []
    bad = [s["code"] for s in manifest["sheets"] if s["qa"] != "PASS"]
    out.append(_check("V-DWG-LAYOUT", "Every sheet passes the automated layout checks (bounds, overlaps, minimum text height)", "FAIL" if bad else "PASS", ", ".join(bad)))
    nums = [s["drawing_number"] for s in manifest["sheets"]]
    out.append(_check("V-DWG-NUMBERS", "Drawing numbers are unique and the sheet count matches the register", "PASS" if len(set(nums)) == len(nums) == len(ds.get("drawings_plan", nums)) else "FAIL", "%d sheets" % len(nums)))
    known = {n["id"] for regs in ds["registers"].values() for n in regs["equipment"]} | {r["id"] for regs in ds["registers"].values() for r in regs["lines"]} | \
            {p["id"] for regs in ds["registers"].values() for p in regs.get("pumps", [])} | {s["id"] for s in ds["structures"]} | {ds["umbilical"]["spec"]["id"]}
    drawn, unknown = set(), set()
    for s in manifest["sheets"]:
        for it in s["items"]:
            drawn.add(it["id"])
            if it["id"] not in known:
                unknown.add(it["id"])
    ref = ds["reference_concept"]
    expected = {n["id"] for n in ds["registers"][ref]["equipment"]} | {r["id"] for r in ds["registers"][ref]["lines"]} | {p["id"] for p in ds["registers"][ref].get("pumps", [])}
    missing = sorted(expected - drawn)
    out.append(_check("V-DWG-ASSETS", "Every model asset and line of the reference concept is drawn, and no drawn identity is unknown to the dataset", "FAIL" if (missing or unknown) else "PASS",
                      "missing: %s; unknown: %s" % (", ".join(missing[:6]) or "none", ", ".join(sorted(unknown)[:6]) or "none")))
    tb = ds["document"]
    unassigned = all(str(tb.get(k, "not assigned")).lower() == "not assigned" for k in ("checked_by", "approved_by")) or bool(ds.get("review"))
    out.append(_check("V-DWG-TITLEBLOCK", "Checker and approver fields stay unassigned until a person is recorded", "PASS" if unassigned else "WARN",
                      "checked_by=%s; approved_by=%s" % (tb.get("checked_by"), tb.get("approved_by"))))
    files_ok = os.path.isfile(manifest["pdf"]) and all(os.path.isfile(s["png"]) and os.path.isfile(s["svg"]) for s in manifest["sheets"])
    out.append(_check("V-DWG-FILES", "SVG, PNG and the vector PDF exist for every sheet", "PASS" if files_ok else "FAIL", "%s; %d sheets" % (os.path.basename(manifest["pdf"]), len(manifest["sheets"]))))
    return out


def _docx_text(path):
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.style.name != "Caption"]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading 1")]
    return "\n".join(parts), heads


def check_document(ds, info, docx_path, pdf_path):
    """Checks on the report: token resolution, bad tokens, mandatory sections, captions, dataset identity and PDF."""
    out = []
    out.append(_check("V-DOC-TOKENS", "Every data token resolves with its unit", "FAIL" if (info["unresolved"] or info["unit_errors"]) else "PASS",
                      "unresolved: %s; unit errors: %s" % (info["unresolved"][:4], info["unit_errors"][:2])))
    text, heads = _docx_text(docx_path)
    bad = sorted({m.group(0) for m in BAD_TOKENS.finditer(text)})
    out.append(_check("V-DOC-BADTOKENS", "No invalid-value marker or unresolved token appears in the document text", "FAIL" if bad else "PASS", ", ".join(bad[:6])))
    missing = [h for h in REQUIRED_HEADINGS if not any(x.startswith(h) for x in heads)]
    out.append(_check("V-DOC-SECTIONS", "The 20 mandatory report chapters are present", "FAIL" if missing else "PASS", ", ".join(missing)))
    out.append(_check("V-DOC-CAPTIONS", "Numbered figures and tables exist", "PASS" if info["n_figures"] > 0 and info["n_tables"] > 0 else "FAIL", "%d figures, %d tables" % (info["n_figures"], info["n_tables"])))
    try:
        with open(info["facts"], encoding="utf-8") as fh:
            facts = json.load(fh)
        same = facts.get("dataset_hash") == ds["meta"]["dataset_hash"]
    except OSError:
        same = False
    out.append(_check("V-DOC-DATASET", "The report was built from the retained dataset (hash agrees)", "PASS" if same else "FAIL", ds["meta"]["dataset_hash"]))
    out.append(_check("V-DOC-OPEN", "Open items, hand-offs and data gaps are visible in the document", "PASS" if ("hand-off" in text.lower() and "data gap" in text.lower()) else "FAIL"))
    if pdf_path is None:
        out.append(_check("V-DOC-PDF", "PDF export", "WARN", "no PDF written (Word or LibreOffice not available or --no-word)"))
    else:
        out.append(_check("V-DOC-PDF", "PDF export exists", "PASS" if os.path.isfile(pdf_path) else "FAIL", pdf_path))
    return out


def compute_readiness(ds, checks):
    """Readiness label from the checks: any FAIL gives NOT_READY; otherwise SCREENING (the tool never assigns REVIEWED)."""
    fails = [c["id"] for c in checks if c["status"] == "FAIL"]
    if fails:
        return {"label": "NOT_READY", "note": "diagnostic package: %d validation check(s) failed" % len(fails),
                "explanation": "Failed checks: %s. The package is a diagnostic and must not be used for design." % ", ".join(fails)}
    reasons = []
    if ds["meta"].get("synthetic"):
        reasons.append("synthetic inputs")
    if ds["handoffs"]:
        reasons.append("%d open specialist hand-off(s)" % len(ds["handoffs"]))
    if ds["data_gaps"]:
        reasons.append("%d data gap(s)" % len(ds["data_gaps"]))
    note = "screening grade: " + ", ".join(reasons) if reasons else "screening grade"
    return {"label": "SCREENING", "note": note, "explanation": "All validation checks passed or carry warnings; the package is screening-grade (%s). Supplier confirmation, specialist analysis, independent "
            "verification and approval are separate states that this tool does not assign." % ", ".join(reasons)}


def summarise(checks, readiness):
    """Counts, overall status and the readiness block."""
    n_fail = sum(1 for c in checks if c["status"] == "FAIL")
    n_warn = sum(1 for c in checks if c["status"] == "WARN")
    return {"status": "FAIL" if n_fail else "PASS", "n_checks": len(checks), "n_pass": sum(1 for c in checks if c["status"] == "PASS"), "n_fail": n_fail, "n_warn": n_warn,
            "readiness": readiness, "checks": checks}


def write_validation(result, out_dir):
    """Write ``validation_report.json`` and ``.md``; returns both paths."""
    os.makedirs(out_dir, exist_ok=True)
    jp, mp = os.path.join(out_dir, "validation_report.json"), os.path.join(out_dir, "validation_report.md")
    with open(jp, "w", encoding="utf-8") as fh:
        json.dump(result, fh, indent=1)
    lines = ["# SURF study validation report", "", "Overall: **%s** (%d checks, %d pass, %d fail, %d warn). Readiness: **%s** - %s" % (
        result["status"], result["n_checks"], result["n_pass"], result["n_fail"], result["n_warn"], result["readiness"]["label"], result["readiness"]["note"]), "",
        "| Check | Result | Description | Detail |", "|---|---|---|---|"]
    for c in result["checks"]:
        lines.append("| %s | %s | %s | %s |" % (c["id"], c["status"], c["description"], str(c["detail"]).replace("|", "/")[:200]))
    with open(mp, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    return jp, mp
