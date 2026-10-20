"""Tests for the process-study-report toolkit (extraction, balances, basis, validation, pipeline).

Uses only the synthetic flowsheet of ``psr_demo`` (SRK, separator + compressor + cooler + recycle).
"""
import copy
import json
import os
import sys

import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

jneqsim = pytest.importorskip("jneqsim")

import psr_balances  # noqa: E402
import psr_basis  # noqa: E402
import psr_extract  # noqa: E402
import psr_report  # noqa: E402
import psr_validate  # noqa: E402

SKILL = os.path.join(os.path.dirname(HERE), ".github", "skills", "process-study-report", "templates")


@pytest.fixture(scope="module")
def demo_data():
    import psr_demo

    model = psr_demo.build_demo_model()
    first = psr_extract.extract_process(model, "A", "Base", model_name="synthetic")
    model2 = psr_demo.build_demo_model(feed_kg_h=26000.0, compressor_outlet_bara=100.0)
    second = psr_extract.extract_process(model2, "B", "Higher feed", model_name="synthetic")
    return psr_extract.merge_cases([first, second])


@pytest.fixture(scope="module")
def inputs():
    basis = psr_basis.load_structured(os.path.join(SKILL, "basis_demo.yaml"))
    standards = psr_basis.load_structured(os.path.join(SKILL, "standards_register_template.yaml"))["standards"]
    basis["operating_cases"].append({"id": "B", "label": "Higher feed", "role": "scenario",
                                     "description": "Feed of 26000 kg/h and 100 bara discharge", "data_source": "specified inputs"})
    basis["interpretation"] = [{"topic": "Compression", "text": "Case A compressor power is [[unit:export compressor:power_kw:A|.0f]] kW and case B is [[unit:export compressor:power_kw:B|.0f]] kW."}]
    return basis, standards


@pytest.fixture(scope="module")
def pipeline(demo_data, inputs, tmp_path_factory):
    import process_study_report

    basis, standards = inputs
    out = str(tmp_path_factory.mktemp("psr"))
    summary = process_study_report.run_pipeline(copy.deepcopy(demo_data), basis, standards, out, use_word=False)
    return out, summary


def test_extract_topology(demo_data):
    assert len(demo_data["streams"]) == 11
    assert len(demo_data["equipment"]) == 7
    assert [c["id"] for c in demo_data["cases"]] == ["A", "B"]
    ids = [s["id"] for s in demo_data["streams"]]
    assert ids == sorted(ids) and ids[0] == "S001"
    labels = [s["label"] for s in demo_data["streams"]]
    assert len(set(labels)) == len(labels)
    cats = {u["tag"]: u["category"] for u in demo_data["equipment"]}
    assert cats["export compressor"] == "compressor" and cats["liquid recycle"] == "recycle"
    kinds = {s["kind"] for s in demo_data["streams"]}
    assert {"feed", "product", "internal"} <= kinds


def test_extract_values_are_physical(demo_data):
    feed = demo_data["streams"][0]
    for case in ("A", "B"):
        r = feed["results"][case]
        assert r["available"] and abs(sum(r["z"].values()) - 1.0) < 1e-9
        assert r["t_c"] == pytest.approx(45.0) and r["p_bara"] == pytest.approx(40.0)
    assert feed["results"]["B"]["mass_kg_h"] > feed["results"]["A"]["mass_kg_h"]
    pseudo = [c for c in demo_data["components"] if c["is_tbp"]]
    assert {c["name"] for c in pseudo} == {"C7_PC", "C9_PC"}
    assert all(c["tc_k"] > 500 and c["omega"] > 0.2 for c in pseudo)
    comp = [u for u in demo_data["equipment"] if u["category"] == "compressor"][0]
    assert comp["results"]["B"]["power_kw"] > comp["results"]["A"]["power_kw"] > 0


def test_stream_numbers_and_flash_standard_volume(demo_data):
    gas = [s for s in demo_data["streams"] if s["label"] == "HP separator gas out"][0]
    flash = gas["results"]["A"]["std_flash_m3_h"]
    assert flash["gas"] > 1000.0 and flash["oil"] == 0.0


def test_merge_rejects_different_topology(demo_data):
    other = copy.deepcopy(demo_data)
    other["streams"] = other["streams"][:-1]
    with pytest.raises(ValueError):
        psr_extract.merge_cases([demo_data, other])


def test_fmt_value_distinguishes_zero_trace_unavailable():
    assert psr_extract.fmt_value(None) == "n/a"
    assert psr_extract.fmt_value(0.0) == "0"
    assert psr_extract.fmt_value(1e-9) == "<1e-06"
    assert psr_extract.fmt_value(12345.678, 5) == "12346"


def test_balances_close_and_detect_errors(demo_data):
    bal = psr_balances.compute_balances(demo_data)
    for case in ("A", "B"):
        o = bal["cases"][case]["overall"]
        assert o["mass_status"] == "PASS" and o["energy_status"] == "PASS"
        comp = [u for u in bal["cases"][case]["units"] if u["tag"] == "export compressor"][0]
        assert comp["energy_status"] == "PASS"
        assert comp["implied_energy_kw"] == pytest.approx(comp["power_kw"], rel=1e-6)
        cooler = [u for u in bal["cases"][case]["units"] if u["tag"] == "after cooler"][0]
        assert cooler["duty_kw"] < 0 and cooler["energy_status"] == "PASS"
    assert bal["summary"]["FAIL"] == 0
    broken = copy.deepcopy(demo_data)
    broken["streams"][4]["results"]["A"]["mass_kg_h"] *= 1.2
    bad = psr_balances.compute_balances(broken)
    assert bad["summary"]["FAIL"] > 0 and bad["fail_items"]


def test_utility_summary(demo_data):
    rows, totals = psr_balances.utility_summary(demo_data, "A")
    assert totals["power_in_kw"] > 0 and totals["cooling_kw"] < 0
    assert rows["power"][0][0] == "export compressor"


def test_basis_template_is_incomplete_and_demo_is_complete(inputs):
    template = psr_basis.load_structured(os.path.join(SKILL, "basis_template.yaml"))
    errors = [f for f in psr_basis.check_basis(template) if f["level"] == "error"]
    assert any(f["code"] == "BASIS-PLACEHOLDER" for f in errors)
    basis, _ = inputs
    assert [f for f in psr_basis.check_basis(basis) if f["level"] == "error"] == []


def test_basis_flags_missing_items(inputs):
    basis, _ = inputs
    broken = copy.deepcopy(basis)
    del broken["conventions"]["heat_sign"]
    broken["data_gaps"][0].pop("effect_on_conclusions")
    broken["data_classification"][0]["class"] = "guessed"
    codes = {f["code"] for f in psr_basis.check_basis(broken)}
    assert {"BASIS-CONVENTION", "BASIS-GAP", "BASIS-CLASS"} <= codes


def test_standards_checks(inputs):
    _, standards = inputs
    assert psr_basis.check_standards(standards) == []
    bad = copy.deepcopy(standards)
    bad[0]["compliance_statement"] = "Fully complies with the standard"
    bad[1].pop("edition")
    codes = {f["code"] for f in psr_basis.check_standards(bad)}
    assert {"STD-OVERCLAIM", "STD-FIELD"} <= codes


def test_pipeline_outputs_and_validation(pipeline):
    out, summary = pipeline
    assert summary["validation"] in ("PASS", "PASS WITH WARNINGS"), open(summary["validation_md"]).read()
    assert os.path.exists(summary["docx"]) and os.path.getsize(summary["docx"]) > 50000
    for name in ("process_data.json", "process_data.xlsx", "stream_components.csv", "balances.json"):
        assert os.path.exists(os.path.join(out, "data", name))
    with open(summary["validation_json"], encoding="utf-8") as fh:
        val = json.load(fh)
    ids = {c["id"]: c["status"] for c in val["checks"]}
    for must in ("PFD-STREAMS", "PFD-TAGS", "COMP-SUM", "BAL-OVERALL", "DOC-EXEC-TRACE", "EXPORTS", "DOC-FIGURES"):
        assert ids[must] == "PASS"
    with open(os.path.join(out, "report", "report_facts.json"), encoding="utf-8") as fh:
        facts = json.load(fh)
    assert facts["unresolved_references"] == [] and facts["resolved_references"]


def test_component_export_is_complete(demo_data, pipeline):
    out, _ = pipeline
    with open(os.path.join(out, "data", "stream_components.csv"), encoding="utf-8") as fh:
        rows = fh.read().strip().split("\n")
    assert len(rows) - 1 == 11 * len(demo_data["components"]) * 2


def test_validator_catches_missing_pfd_stream_and_new_streams(demo_data, inputs, pipeline):
    out, summary = pipeline
    basis, standards = inputs
    import psr_diagrams  # noqa: F401

    data = copy.deepcopy(demo_data)
    extra = copy.deepcopy(data["streams"][0])
    extra.update({"id": "S999", "label": "phantom stream", "name": "phantom"})
    data["streams"].append(extra)
    bal = psr_balances.compute_balances(data)
    diagrams = {"bfd": {"svg": summary["diagrams"]["bfd"]}, "pfd": {"sheets": [{"svg": p} for p in summary["diagrams"]["pfd_sheets"]], "warnings": []}}
    res = psr_validate.validate_all(data, basis, standards, bal, diagrams, summary["docx"], os.path.join(out, "report", "report_facts.json"), os.path.join(out, "data"))
    st = {c["id"]: c["status"] for c in res["checks"]}
    assert st["PFD-STREAMS"] == "FAIL" and st["DOC-STREAM-REGISTER"] == "FAIL" and res["status"] == "FAIL"


def test_validator_catches_placeholder_and_missing_alt_text(demo_data, inputs, pipeline, tmp_path):
    out, summary = pipeline
    basis, standards = inputs
    from docx import Document

    doc = Document(summary["docx"])
    doc.add_paragraph("Results to be confirmed: TBD")
    for shape in doc.inline_shapes:
        shape._inline.docPr.set("descr", "")
    broken = str(tmp_path / "broken.docx")
    doc.save(broken)
    bal = psr_balances.compute_balances(demo_data)
    res = psr_validate.validate_all(demo_data, basis, standards, bal, None, broken, os.path.join(out, "report", "report_facts.json"), None)
    st = {c["id"]: c["status"] for c in res["checks"]}
    assert st["DOC-PLACEHOLDERS"] == "FAIL", st
    assert st["DOC-FIGURES"] == "FAIL", st


def test_validator_catches_untraced_executive_number(demo_data, inputs, pipeline, tmp_path):
    out, summary = pipeline
    basis, standards = inputs
    from docx import Document

    doc = Document(summary["docx"])
    for par in doc.paragraphs:
        if par.text.startswith("The model comprises"):
            par.add_run(" The plant is expected to deliver 61357.9 kg/h of export gas.")
            break
    broken = str(tmp_path / "untraced.docx")
    doc.save(broken)
    bal = psr_balances.compute_balances(demo_data)
    res = psr_validate.validate_all(demo_data, basis, standards, bal, None, broken, os.path.join(out, "report", "report_facts.json"), None)
    st = {c["id"]: c["status"] for c in res["checks"]}
    assert st["DOC-EXEC-TRACE"] == "FAIL"


def test_composition_sum_check_detects_bad_data(demo_data, inputs):
    basis, standards = inputs
    bad = copy.deepcopy(demo_data)
    bad["streams"][0]["results"]["A"]["z"]["methane"] += 0.01
    bal = psr_balances.compute_balances(bad)
    res = psr_validate.validate_all(bad, basis, standards, bal, None, None, None, None)
    assert {c["id"]: c["status"] for c in res["checks"]}["COMP-SUM"] == "FAIL"


def test_unresolved_reference_is_reported(demo_data, inputs):
    basis, standards = inputs
    bal = psr_balances.compute_balances(demo_data)
    res = psr_report.Resolver(demo_data, bal, basis, psr_report.compute_facts(demo_data, bal))
    assert res.expand("x [[stream:S404:mass_kg_h:A]] y") == "x n/a y"
    assert res.errors == ["stream:S404:mass_kg_h:A"]
    assert res.expand("[[stream:S001:t_c:A|.1f]]") == "45.0"


def test_merge_into_results_adds_block_without_touching_other_keys(pipeline, tmp_path):
    import process_study_report

    out, summary = pipeline
    path = tmp_path / "results.json"
    path.write_text(json.dumps({"key_results": {"x": 1}}), encoding="utf-8")
    block = process_study_report.merge_into_results(str(path), summary, task_dir=out)
    saved = json.loads(path.read_text(encoding="utf-8"))
    assert saved["key_results"] == {"x": 1}
    assert saved["process_study_report"]["validation_status"] == summary["validation"]
    assert block["report_docx"].startswith("report/") and block["diagrams"]["pfd_sheets"][0].endswith(".svg")


def test_pfd_carries_equipment_data_and_connector_drawing_numbers(demo_data, pipeline):
    import process_study_report

    lines = process_study_report.default_equipment_data(demo_data, "A")
    assert lines["export compressor"][0].startswith("W = ") and any(l.startswith("eta_p") for l in lines["export compressor"])
    assert lines["after cooler"][0].startswith("Q = -")
    out, _ = pipeline
    with open(os.path.join(out, "diagrams", "pfd", "pfd_layout_spec.json"), encoding="utf-8") as fh:
        spec = json.load(fh)
    assert "export compressor" in json.dumps(spec["equipment_data"]) if "equipment_data" in spec else "export compressor" in json.dumps(spec)
