"""Report fidelity, evidence labels and parameter completeness for PVT studies."""

import copy
import json
import math
import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent))
import pvt_report  # noqa: E402


@pytest.fixture
def study():
    """A synthetic numerical fixture; no result is called laboratory evidence."""
    components = [{"name": "methane", "moleFraction": 0.7, "criticalTemperatureK": 190.56},
                  {"name": "n-decane", "moleFraction": 0.3, "criticalTemperatureK": 617.7}]
    parameters = {"systemClass": "neqsim.thermo.system.SystemSrkEos", "components": components,
                  "phases": [{"index": 0, "type": "gas", "binaryInteractionParameters": [
                      {"i": i, "j": j, "kij": 0.025 if i != j else 0.0, "kijT1": 0.0}
                      for i in range(2) for j in range(2)]}]}
    experiment = {"id": "CCE-1", "sample_id": "S1", "model_id": "FINAL", "kind": "CCE",
                  "temperature_k": 350.0, "pressures_bara": [300.0, 200.0, 100.0],
                  "role": "calibration", "source_kind": "synthetic_reference",
                  "predictions": {"values": {"relativeVolume": [0.98, 1.0, 1.3]}}}
    return {"study": {"id": "PVT-DEMO", "title": "PVT fluid characterization study",
                       "revision": "A", "readiness": "NUMERICALLY VERIFIED", "synthetic": True},
            "samples": [{"id": "S1", "synthetic": True, "origin": "Synthetic workflow fixture"}],
            "analyses": [{"id": "A1", "sample_id": "S1", "composition_basis": "mole"}],
            "models": [{"id": "FINAL", "stage": "final", "sample_id": "S1", "capture": parameters}],
            "parameters": parameters, "experiments": [experiment],
            "observations": [{"id": "O1", "experiment_id": "CCE-1", "property": "relative_volume",
                              "unit": "m3/m3 at saturation", "values": [0.99, 1.0, 1.31],
                              "predicted_values": [0.98, 1.0, 1.3], "role": "calibration",
                              "uncertainty": [0.01, 0.01, 0.02]}],
            "consistency": [{"model_id": "FINAL", "checks": {"compositionSymmetry": True}, "status": "PASS"}],
            "exports": [{"id": "E1", "model_id": "FINAL", "status": "PASS",
                         "audit": {"targetSimulator": "UNVERIFIED"}}],
            "findings": [{"code": "SYNTHETIC", "status": "WARN", "message": "No laboratory validation"}]}


def test_render_preserves_dataset_and_distinguishes_numerical_evidence(study, tmp_path):
    before = copy.deepcopy(study)
    result = pvt_report.build_report(study, tmp_path)
    assert study == before
    assert Path(result["docx"]).is_file()
    assert Path(result["html"]).is_file()
    doc = Document(result["docx"])
    prose = "\n".join(p.text for p in doc.paragraphs)
    assert "SYNTHETIC WORKFLOW DEMONSTRATION" in prose
    assert "Numerical replay does not establish experimental accuracy" in prose
    assert "Target-simulator compatibility remains unverified" in prose
    assert result["physical_acceptance"] == "NOT INFERRED"
    assert "REQUIRED" in result["visual_inspection"]
    if result["pdf"]:
        import fitz

        with fitz.open(result["pdf"]) as pdf:
            text = "".join(page.get_text() for page in pdf)
            assert "16 Parameter appendices and reproducibility" in text
            assert "No laboratory validation" in text
            assert len(pdf) >= 3


def test_complete_zero_bip_entries_and_component_parameters_are_in_appendix(study):
    blocks = pvt_report.report_blocks(study, [])
    tables = [b for b in blocks if b["type"] == "table"]
    properties = next(b for b in tables if "full component properties" in b["title"])
    assert ["criticalTemperatureK", "190.56", "617.7"] in properties["rows"]
    matrix = next(b for b in tables if "kij matrix" in b["title"])
    assert matrix["rows"] == [["0", "0", "0.025"], ["1", "0.025", "0"]]
    temperature_matrix = next(b for b in tables if "kijT1 matrix" in b["title"])
    assert all(r[1:] == ["0", "0"] for r in temperature_matrix["rows"])


def test_computed_failure_status_cannot_be_overridden_by_input_readiness(study, tmp_path):
    study["status"] = "FAIL"
    study["study"]["readiness"] = "ACCEPTED"
    result = pvt_report.build_report(study, tmp_path)
    content = Path(result["html"]).read_text(encoding="utf-8")
    assert "The study readiness is FAIL" in content
    doc = Document(result["docx"])
    assert any("Status FAIL" in paragraph.text for paragraph in doc.paragraphs)


def test_calculated_recombination_closure_assumptions_and_gaps_are_reported(study):
    study["recombination"] = [{"id": "R1", "basis": "molar_amount"}]
    study["recombination_audit"] = [{"id": "R1", "mole_balance_error": 0.0, "status": "PASS"}]
    study["assumptions"] = ["Declared recombination ratio"]
    study["data_gaps"] = ["Sampling conditions unavailable"]
    tables = {block["title"]: block for block in pvt_report.report_blocks(study, [])
              if block["type"] == "table"}
    assert ("recombination_audit[0].mole_balance_error", "0") in tables["Calculated recombination audit"]["rows"]
    assert ("assumptions[0]", "Declared recombination ratio") in tables["Assumptions"]["rows"]
    assert ("data_gaps[0]", "Sampling conditions unavailable") in tables["Data gaps"]["rows"]


def test_incomplete_ordered_bip_register_cannot_render_as_complete(study):
    study["parameters"]["phases"][0]["binaryInteractionParameters"].pop()
    with pytest.raises(ValueError, match="every ordered pair"):
        pvt_report.report_blocks(study, [])


def test_observations_use_canonical_predicted_arrays_and_roles(study):
    study["observations"][0]["role"] = "validation"
    study["observations"][0]["predicted_values"] = [1.1, 1.2, 1.3]
    rows = pvt_report._observations(study)
    assert len(rows) == 3
    assert [r["predicted"] for r in rows] == [1.1, 1.2, 1.3]
    assert [r["measured"] for r in rows] == [0.99, 1.0, 1.31]
    assert all(r["role"] == "validation" for r in rows)
    assert rows[0]["pressure_bara"] == 300.0


def test_unresolved_tokens_block_creation(study, tmp_path):
    study["study"]["objective"] = "Check [[missing:fluid:pressure]]"
    with pytest.raises(ValueError, match="Unresolved report token"):
        pvt_report.build_report(study, tmp_path / "not-created")
    assert not (tmp_path / "not-created").exists()


def test_html_escapes_external_evidence_and_uses_safe_filename(study, tmp_path):
    study["study"]["id"] = "../../PVT<script>"
    study["samples"][0]["origin"] = "<script>alert('bad')</script>"
    result = pvt_report.build_report(study, tmp_path)
    assert Path(result["html"]).parent == tmp_path
    content = Path(result["html"]).read_text(encoding="utf-8")
    assert "<script>" not in content
    assert "&lt;script&gt;" in content
    manifest = json.loads((tmp_path / "report_manifest.json").read_text())
    assert len(manifest["dataset_sha256"]) == 64


def test_only_registered_protocols_are_drawn_and_svg_text_remains_editable(study, tmp_path):
    figures = pvt_report.make_figures(study, tmp_path)
    names = {f["id"] for f in figures}
    assert "protocol_cce" in names
    assert not {"protocol_cvd", "protocol_dle", "protocol_separator"} & names
    svg = (tmp_path / "protocol_cce.svg").read_text(encoding="utf-8")
    assert "<text" in svg
    assert "gas and liquid retained" in svg
    assert "350" in svg
    assert "300" in svg
    assert all(Path(f["pdf"]).is_file() for f in figures)


def test_missing_evidence_is_unavailable_and_never_inferred_as_pass(tmp_path):
    result = pvt_report.build_report({"study": {"id": "EMPTY"}}, tmp_path)
    content = Path(result["html"]).read_text(encoding="utf-8")
    assert "Unavailable - evidence not supplied" in content
    assert "UNASSESSED" in content
    assert result["physical_acceptance"] == "NOT INFERRED"


def test_legacy_report_text_cannot_create_an_unsplittable_table_row(study, tmp_path):
    study["experiments"][0]["predictions"]["legacyReportMarkdown"] = "Legacy report line\n" * 1000
    result = pvt_report.build_report(study, tmp_path)
    content = Path(result["html"]).read_text(encoding="utf-8")
    assert "Legacy text retained in the machine-readable dataset" in content
    assert "Legacy report line" not in content


def test_separator_sketch_uses_actual_stage_conditions_not_reservoir_pressure(study, tmp_path):
    experiment = study["experiments"][0]
    experiment["kind"] = "SEPARATOR"
    experiment["stages"] = [
        {"pressure": {"value": 40, "unit": "bara"}, "temperature": {"value": 50, "unit": "C"}},
        {"pressure": {"value": 1.01325, "unit": "bara"}, "temperature": {"value": 15, "unit": "C"}},
    ]
    pvt_report.make_figures(study, tmp_path)
    svg = (tmp_path / "protocol_separator.svg").read_text(encoding="utf-8")
    assert "P = 40 bara" in svg
    assert "T = 323.15 K" in svg
    assert "P = 40 to 1.01325 bara; 2 pressure steps" in svg


def _viscosity_study(values, availability):
    """A pointwise viscosity fixture whose availability is independent of value."""
    return {"study": {"id": "VISCOSITY-QA"}, "experiments": [
        {"id": "VISC-1", "kind": "VISCOSITY", "pressures_bara": [300, 200, 100],
         "predictions": {"values": {"aqueousViscosityPaS": values},
                         "phaseAvailability": {"aqueous": availability}}}]}


def test_absent_aqueous_viscosity_is_explicitly_unavailable_in_report(tmp_path, monkeypatch):
    dataset = _viscosity_study([None, None, None], [False, False, False])
    before = copy.deepcopy(dataset)
    captured = []
    axes = []
    save = pvt_report._save_figure

    def inspect_figure(fig, directory, stem, caption, source):
        if stem.startswith("prediction_"):
            captured.append([line.get_ydata().tolist() for line in fig.axes[0].lines])
            axes.append((fig.axes[0].get_xlim(), fig.axes[0].get_yticks().tolist()))
        return save(fig, directory, stem, caption, source)

    monkeypatch.setattr(pvt_report, "_save_figure", inspect_figure)
    report = pvt_report.build_report(dataset, tmp_path)
    assert captured == [[]]
    assert axes == [((100, 300), [])]
    assert dataset == before
    svg = (tmp_path / "figures" / "prediction_000.svg").read_text(encoding="utf-8")
    assert "Unavailable - aqueous phase absent at all registered points" in svg
    content = Path(report["html"]).read_text(encoding="utf-8")
    assert "Unavailable - aqueous phase absent" in content
    assert "phaseAvailability.aqueous[0]" in content


def test_phase_disappearance_leaves_a_plot_gap_and_preserves_available_zero(tmp_path, monkeypatch):
    dataset = _viscosity_study([0.0, None, 0.001], [True, False, True])
    captured = []
    markers = []
    save = pvt_report._save_figure

    def inspect_figure(fig, directory, stem, caption, source):
        if stem.startswith("prediction_"):
            captured.append(fig.axes[0].lines[0].get_ydata().tolist())
            markers.append(fig.axes[0].lines[0].get_marker())
        return save(fig, directory, stem, caption, source)

    monkeypatch.setattr(pvt_report, "_save_figure", inspect_figure)
    pvt_report.make_figures(dataset, tmp_path)
    assert len(captured) == 1
    assert captured[0][0] == 0.0
    assert math.isnan(captured[0][1])
    assert captured[0][2] == 0.001
    assert markers == ["o"]
    svg = (tmp_path / "prediction_000.svg").read_text(encoding="utf-8")
    assert "aqueous phase absent at 1 registered point(s); curve gaps retained" in svg
    rows = pvt_report._observations(dataset)
    assert [row["predicted"] for row in rows] == [0.0, None, 0.001]


def test_explicit_phase_availability_overrides_old_absence_encoding_without_guessing():
    dataset = _viscosity_study([0.0, 0.0, 0.001], [True, False, True])
    assert [row["predicted"] for row in pvt_report._observations(dataset)] == [0.0, None, 0.001]
    dataset["experiments"][0]["predictions"].pop("phaseAvailability")
    assert [row["predicted"] for row in pvt_report._observations(dataset)] == [0.0, 0.0, 0.001]


@pytest.mark.parametrize("mask", [[False], [True, 0, False], [True, None, False], False, None, "false"])
def test_malformed_explicit_phase_mask_cannot_render_a_physical_curve(mask, tmp_path):
    dataset = _viscosity_study([0.0, 0.0, 0.001], mask)
    with pytest.raises(ValueError, match="phaseAvailability"):
        pvt_report.build_report(dataset, tmp_path / "rejected")
    assert not (tmp_path / "rejected").exists()


def test_phase_mask_must_match_registered_pressure_grid(tmp_path):
    dataset = _viscosity_study([None, 0.001], [False, True])
    with pytest.raises(ValueError, match="registered pressures"):
        pvt_report.make_figures(dataset, tmp_path / "rejected")


@pytest.mark.parametrize("masks", [None, [], False])
def test_explicit_phase_availability_metadata_must_be_a_dictionary(masks):
    dataset = _viscosity_study([None, None, None], [False, False, False])
    dataset["experiments"][0]["predictions"]["phaseAvailability"] = masks
    with pytest.raises(ValueError, match="phaseAvailability must be a dictionary"):
        pvt_report._observations(dataset)


def test_explicit_phase_availability_cannot_omit_a_reported_viscosity_phase():
    dataset = _viscosity_study([0.0, 0.001, 0.002], [True, True, True])
    dataset["experiments"][0]["predictions"]["phaseAvailability"] = {}
    with pytest.raises(ValueError, match="phaseAvailability.aqueous is required"):
        pvt_report._observations(dataset)


@pytest.mark.parametrize("values, availability", [
    ([None, 0.001, 0.002], [False, True, True]),
    ([0.001, 0.002, None], [True, True, False]),
])
def test_absent_endpoint_remains_on_pressure_axis_and_annotation_avoids_data(values, availability,
                                                                          tmp_path, monkeypatch):
    dataset = _viscosity_study(values, availability)
    captured = []
    save = pvt_report._save_figure

    def inspect_figure(fig, directory, stem, caption, source):
        if stem.startswith("prediction_"):
            ax = fig.axes[0]
            annotation = next(text for text in ax.texts if "phase absent" in text.get_text())
            captured.append((ax.get_xlim(), annotation.get_position()[1]))
        return save(fig, directory, stem, caption, source)

    monkeypatch.setattr(pvt_report, "_save_figure", inspect_figure)
    pvt_report.make_figures(dataset, tmp_path)
    assert len(captured) == 1
    assert captured[0][0] == (100, 300)
    assert captured[0][1] < 0
