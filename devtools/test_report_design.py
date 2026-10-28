"""Regression tests for template-preserving report finishing and visual QA."""
import hashlib
import json
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import report_design as design


def _report(path):
    """Make a synthetic report with branded front matter and varied body tables."""
    doc = Document()
    doc.styles["Normal"].font.name = "Garamond"
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Heading 1"].font.color.rgb = RGBColor.from_string("265080")
    doc.sections[0].header.paragraphs[0].text = "Synthetic engineering brand"
    doc.sections[0].footer.paragraphs[0].text = "Internal - unreviewed"
    doc.add_paragraph("Synthetic cover")
    cover = doc.add_table(rows=1, cols=2)
    cover.cell(0, 0).text = "Revision"
    cover.cell(0, 1).text = "A"
    doc.add_heading("Results", 1)
    paragraph = doc.add_paragraph("Flow is 12.3 kg/h.")
    paragraph.runs[0].font.size = Pt(12)
    paragraph.add_run(" Re = 100.").italic = True
    table = doc.add_table(rows=4, cols=3)
    table.columns[0].width = Inches(2)
    table.columns[1].width = Inches(1)
    table.columns[2].width = Inches(3)
    for cell, text in zip(table.rows[0].cells, ["Quantity", "Value", "Evidence"]):
        cell.text = text
    for row in table.rows[1:]:
        for cell, text in zip(row.cells, ["Flow", "12.3", "Independent check"]):
            cell.text = text
    table.cell(2, 0).merge(table.cell(2, 1)).text = "Merged evidence"
    doc.add_heading("Validation", 2)
    caption = doc.add_paragraph(style="Caption")
    caption.add_run("Table ")
    design._field(caption, r"SEQ Table \* ARABIC")
    caption.add_run(": Existing validation caption")
    second = doc.add_table(rows=2, cols=2)
    second.cell(0, 0).text = "Check"
    second.cell(0, 1).text = "Result"
    second.cell(1, 0).text = "Mass balance"
    second.cell(1, 1).text = "0.01 %"
    doc.save(path)
    return path


def test_finishing_preserves_source_content_brand_and_special_runs(tmp_path):
    source = _report(tmp_path / "source.docx")
    before = source.read_bytes()
    output = tmp_path / "pilot.docx"
    result = design.finish_document(source, output)
    assert source.read_bytes() == before
    assert result["source_sha256"] == hashlib.sha256(before).hexdigest()
    assert result["original_content_preserved"] is True
    assert result["body_tables"] == 2
    assert result["added_table_captions"] == ["Results"]
    assert result["visual_review"] == "PENDING"
    doc = Document(output)
    assert doc.styles["Normal"].font.name == "Garamond"
    assert str(doc.styles["Heading 1"].font.color.rgb) == "265080"
    assert doc.sections[0].header.paragraphs[0].text == "Synthetic engineering brand"
    assert "Internal - unreviewed" in doc.sections[0].footer.paragraphs[0].text
    body = next(p for p in doc.paragraphs if p.text.startswith("Flow is"))
    assert body.runs[0].font.size is None
    assert body.runs[1].italic
    assert doc.tables[0].cell(0, 0).text == "Revision"
    assert not doc.tables[0]._tbl.xpath(".//w:tblHeader")


def test_body_tables_have_consistent_widths_headers_and_merged_cells(tmp_path):
    source = _report(tmp_path / "source.docx")
    output = tmp_path / "pilot.docx"
    design.finish_document(source, output)
    doc = Document(output)
    for table in doc.tables[1:]:
        section = doc.sections[0]
        width = section.page_width - section.left_margin - section.right_margin
        assert abs(sum(c.width for c in table.columns) - width) < 635 * 3
        assert len(table._tbl.xpath(".//w:tblHeader")) == 1
        assert len(table._tbl.xpath(".//w:cantSplit")) == len(table.rows)
        assert table.cell(0, 0).paragraphs[0].paragraph_format.keep_with_next
        assert table.cell(1, 1).paragraphs[0].alignment == design.WD_ALIGN_PARAGRAPH.RIGHT
        assert all(run.font.size == Pt(10.5) for row in table.rows
                   for cell in row.cells for p in cell.paragraphs for run in p.runs)
    merged = doc.tables[1].cell(2, 0)
    expected = doc.tables[1].columns[0].width + doc.tables[1].columns[1].width
    assert abs(merged.width - expected) < 635 * 3


def test_finishing_is_idempotent_for_captions_and_footer_fields(tmp_path):
    source = _report(tmp_path / "source.docx")
    first, second = tmp_path / "first.docx", tmp_path / "second.docx"
    design.finish_document(source, first)
    result = design.finish_document(first, second)
    assert result["added_table_captions"] == []
    doc = Document(second)
    assert len([p for p in doc.paragraphs if design._is_caption(p._p)]) == 2
    instructions = design._instructions(doc.sections[0].footer._element)
    assert len([text for text in instructions if text.strip() == "PAGE"]) == 1


def test_section_geometry_and_oversized_rows(tmp_path):
    source = _report(tmp_path / "source.docx")
    doc = Document(source)
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    doc.add_heading("Appendix", 1)
    table = doc.add_table(rows=2, cols=2)
    table.cell(1, 0).text = "Very long evidence. " * 120
    doc.save(source)
    output = tmp_path / "pilot.docx"
    result = design.finish_document(source, output, table_label="Tabell")
    doc = Document(output)
    last = doc.tables[-1]
    width = doc.sections[-1].page_width - Inches(1)
    assert abs(sum(c.width for c in last.columns) - width) < 635 * 3
    assert not last.rows[1]._tr.xpath("./w:trPr/w:cantSplit")
    assert result["warnings"]
    assert any(p.text.startswith("Tabell ") for p in doc.paragraphs)


def test_short_tables_stay_together_and_long_evidence_gets_more_width(tmp_path):
    source = _report(tmp_path / "source.docx")
    doc = Document(source)
    doc.tables[-1].cell(1, 1).text = "Independent evidence detail. " * 25
    doc.save(source)
    output = tmp_path / "pilot.docx"
    design.finish_document(source, output)
    table = Document(output).tables[-1]
    assert table.columns[1].width > 2 * table.columns[0].width
    assert table.rows[0].cells[0].paragraphs[0].paragraph_format.keep_with_next
    assert not table.rows[-1].cells[0].paragraphs[0].paragraph_format.keep_with_next


def test_output_guards_and_unrecognized_documents(tmp_path):
    source = _report(tmp_path / "source.docx")
    with pytest.raises(ValueError, match="different files"):
        design.finish_document(source, source, overwrite=True)
    output = tmp_path / "exists.docx"
    output.write_bytes(b"existing")
    with pytest.raises(FileExistsError):
        design.finish_document(source, output)
    assert output.read_bytes() == b"existing"
    unknown = tmp_path / "unknown.docx"
    Document().save(unknown)
    with pytest.raises(ValueError, match="No body heading"):
        design.finish_document(unknown, tmp_path / "new.docx")
    with pytest.raises(ValueError, match="extension"):
        design.finish_document(source, tmp_path / "bad.pdf")


def test_pdf_preview_has_every_page_and_does_not_claim_visual_approval(tmp_path):
    fitz = pytest.importorskip("fitz")
    pdf = tmp_path / "test.pdf"
    with fitz.open() as doc:
        for _ in range(3):
            page = doc.new_page()
            page.insert_text((72, 72), "Synthetic report - 12.3 kg/h")
        doc.save(pdf)
    result = design.preview_pdf(pdf, tmp_path / "preview")
    assert result["pages"] == 3
    assert all(Path(path).exists() for path in result["page_images"])
    assert Path(result["contact_sheet"]).exists()
    assert result["geometry_findings"] == []
    assert result["visual_review"] == "PENDING"
    shorter = tmp_path / "shorter.pdf"
    with fitz.open() as doc:
        doc.new_page().insert_text((72, 72), "A shorter revision")
        doc.save(shorter)
    directory = tmp_path / "preview"
    (directory / "review_notes.png").write_bytes(b"reviewer-owned")
    design.preview_pdf(shorter, directory)
    assert len(list(directory.glob("page_*.png"))) == 1
    assert (directory / "review_notes.png").read_bytes() == b"reviewer-owned"


def test_cli_records_render_failure_explicitly(tmp_path, monkeypatch, capsys):
    source = _report(tmp_path / "source.docx")
    output = tmp_path / "pilot.docx"

    def fail(*args):
        raise RuntimeError("Renderer unavailable")

    monkeypatch.setattr(design, "render_pdf", fail)
    assert design.main([str(source), "--output", str(output), "--render"]) == 2
    assert "Renderer unavailable" in capsys.readouterr().err
    qa = json.loads(output.with_suffix(".qa.json").read_text())
    assert qa["render_error"] == "Renderer unavailable"
    assert qa["visual_review"] == "PENDING"


def test_missing_renderer_is_not_a_success(tmp_path, monkeypatch):
    monkeypatch.setattr(sys, "platform", "linux")
    monkeypatch.setattr(design.shutil, "which", lambda name: None)
    with pytest.raises(RuntimeError, match="no PDF was rendered"):
        design.render_pdf(tmp_path / "input.docx", tmp_path / "output.pdf")


def test_cli_refuses_to_replace_existing_output_companions(tmp_path, capsys):
    source = _report(tmp_path / "source.docx")
    output = tmp_path / "pilot.docx"
    pdf = output.with_suffix(".pdf")
    pdf.write_bytes(b"existing approved PDF")
    assert design.main([str(source), "--output", str(output), "--render"]) == 2
    assert "companion already exists" in capsys.readouterr().err
    assert pdf.read_bytes() == b"existing approved PDF"
    assert not output.exists()


def test_cli_dispatch_preserves_callers_relative_paths(monkeypatch, tmp_path):
    import neqsim_cli
    seen = []
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(design, "main", lambda args: seen.extend(args) or 0)
    monkeypatch.setattr(sys, "argv",
                        ["neqsim", "report-design", "source.docx", "--output", "copy.docx"])
    with pytest.raises(SystemExit) as result:
        neqsim_cli.main()
    assert result.value.code == 0
    assert seen == ["source.docx", "--output", "copy.docx"]
    assert Path.cwd() == tmp_path
