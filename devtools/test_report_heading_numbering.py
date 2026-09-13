"""Report heading-numbering tests (devtools/task_template/step3_report)."""
import importlib.util
import os
import sys

import pytest

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GENERATOR = os.path.join(REPO, "devtools", "task_template", "step3_report",
                         "generate_report.py")
docx = pytest.importorskip("docx")


@pytest.fixture(scope="module")
def gen(tmp_path_factory):
    """Import the generator module bound to a throwaway task folder."""
    task = tmp_path_factory.mktemp("task")
    (task / "step3_report").mkdir()
    argv = sys.argv
    sys.argv = ["generate_report.py", "--task-dir", str(task)]
    try:
        spec = importlib.util.spec_from_file_location("neqsim_report_gen", GENERATOR)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
    finally:
        sys.argv = argv
    return module


def _numbered_heading_style(document):
    """Give Heading 1 automatic list numbering, as corporate templates do."""
    from docx.oxml.ns import nsdecls, qn
    from docx.oxml import parse_xml

    style = document.styles["Heading 1"].element
    p_pr = style.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = parse_xml("<w:pPr {}/>".format(nsdecls("w")))
        style.append(p_pr)
    p_pr.append(parse_xml(
        '<w:numPr {}><w:ilvl w:val="0"/><w:numId w:val="3"/></w:numPr>'.format(
            nsdecls("w"))))


def test_manual_number_kept_when_template_does_not_number(gen):
    document = docx.Document()
    gen._HEADING_NUMBERING_CACHE.clear()
    heading = gen._add_heading(document, "7. Solution Workflow", level=1)
    assert heading.text == "7. Solution Workflow"


def test_manual_number_dropped_when_template_numbers_headings(gen):
    document = docx.Document()
    _numbered_heading_style(document)
    gen._HEADING_NUMBERING_CACHE.clear()
    heading = gen._add_heading(document, "7. Solution Workflow", level=1)
    # Word supplies the number, so only one numbering scheme survives.
    assert heading.text == "Solution Workflow"


def test_front_matter_heading_is_unnumbered(gen):
    from docx.oxml.ns import qn

    document = docx.Document()
    _numbered_heading_style(document)
    gen._HEADING_NUMBERING_CACHE.clear()
    heading = gen._add_heading(document, "Table of Contents", level=1,
                               numbered=False)
    num_pr = heading._p.find(qn("w:pPr")).find(qn("w:numPr"))
    assert num_pr.find(qn("w:numId")).get(qn("w:val")) == "0"


def test_sections_are_renumbered_without_gaps(gen):
    sections = [
        {"heading": "1. Executive Summary"},
        {"heading": "3. Results"},          # counter bug upstream
        {"heading": "3. Discussion"},       # duplicate
        {"heading": "Conclusions"},         # never numbered
    ]
    renumbered = [item["heading"] for item in gen._renumber_sections(sections)]
    assert renumbered == ["1. Executive Summary", "2. Results",
                          "3. Discussion", "4. Conclusions"]
