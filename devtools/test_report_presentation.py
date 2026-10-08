"""End-to-end tests for figure placement, cross-references, results presentation,
equations, metadata and offline HTML in the task report generator."""
import json
import re
import shutil
import struct
import subprocess
import sys
import zipfile
import zlib
from pathlib import Path

import pytest

GENERATOR = Path(__file__).resolve().parent / "task_template" / "step3_report" / "generate_report.py"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


def _png():
    def chunk(tag, data):
        return (struct.pack(">I", len(data)) + tag + data
                + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF))
    raw = b"\x00\x00\x00\x00"
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def _task(root, results):
    task = root / "2026-10-07_presentation_task"
    (task / "step3_report").mkdir(parents=True)
    (task / "step1_scope_and_research").mkdir()
    (task / "figures").mkdir()
    for name in ("fig01_overview.png", "fig02_pressure.png"):
        (task / "figures" / name).write_bytes(_png())
    shutil.copy(str(GENERATOR), str(task / "step3_report" / "generate_report.py"))
    shutil.copy(str(GENERATOR.with_name("report_kit.py")), str(task / "step3_report"))
    shutil.copytree(str(GENERATOR.with_name("vendor")), str(task / "step3_report" / "vendor"))
    (task / "results.json").write_text(json.dumps(results), encoding="utf-8")
    (task / "study_config.yaml").write_text(
        'study:\n  title: "Presentation Task"\n  author: "Test Author"\nreport:\n'
        '  formats:\n    - docx\n    - html\n', encoding="utf-8")
    return task


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    pytest.importorskip("docx")
    results = {
        "key_results": {
            "outlet_temperature": {"value": -18.456, "unit": "C", "headline": True,
                                    "label": "Outlet temperature", "decimals": 1,
                                    "group": "Thermal"},
            "pressure_drop_bar": {"value": 3.2, "headline": True},
            "raw_internal_flag_a": 1.5, "raw_internal_flag_b": 2.5,
        },
        "approach": "Used SRK. See {fig:fig02_pressure} for the profile.",
        "conclusions": "Outlet is -18.5 C and the drop is 3.2 bar.",
        "equations": [
            {"label": "Reynolds number", "latex": r"Re = \frac{\rho v D}{\mu}"},
            {"label": "Piecewise", "latex": r"f = \begin{cases} 1 & x<0 \\ 2 & x>0 \end{cases}"},
        ],
        "figure_captions": {"fig01_overview.png": "Overview", "fig02_pressure.png": "Pressure"},
        "figure_discussion": [
            {"figure": "fig02_pressure.png", "title": "Pressure profile",
             "observation": "Drop of 3.2 bar, see {fig:fig02_pressure}.",
             "linked_results": ["pressure_drop_bar", "outlet_temperature"]},
            {"figure": "fig02_pressure.png", "title": "Second look",
             "observation": "Same figure again."},
        ],
        "references": [{"id": "R1", "text": "Smith (2019)."}],
    }
    task = _task(tmp_path_factory.mktemp("pres"), results)
    run = subprocess.run([sys.executable, str(task / "step3_report" / "generate_report.py")],
                         cwd=str(task), capture_output=True, text=True)
    assert run.returncode == 0, run.stdout + run.stderr
    report = task / "step3_report"
    return {"docx": next(report.glob("*.docx")), "html": next(report.glob("*.html")).read_text(
        encoding="utf-8")}


def _document_xml(path):
    return zipfile.ZipFile(str(path)).read("word/document.xml").decode("utf-8")


def test_discussed_figure_sits_in_discussion_and_only_once(built):
    xml = _document_xml(built["docx"])
    assert xml.count("<w:drawing") == 2  # overview in Results, pressure in Discussion
    assert re.search(r"Discussion 1: Pressure profile", xml)
    assert xml.index("Discussion 1: Pressure profile") < xml.index('descr="Pressure')
    assert "See" in xml  # second entry points back at the figure


def test_word_cross_references_are_live_fields_to_bookmarks(built):
    xml = _document_xml(built["docx"])
    assert xml.count("w:bookmarkStart") >= 2
    assert re.search(r"instrText[^>]*> REF _RefFig_fig02_pressure \\h", xml)
    assert 'w:name="_RefFig_fig02_pressure"' in xml


def test_figures_have_alt_text_and_document_has_metadata(built):
    xml = _document_xml(built["docx"])
    assert len(re.findall(r'<wp:docPr[^>]*descr="[^"]+"', xml)) == 2
    import docx
    props = docx.Document(str(built["docx"])).core_properties
    assert props.title == "Presentation Task"
    assert props.author == "Test Author"
    assert "NeqSim" in props.keywords


def test_headline_results_lead_and_full_list_moves_to_appendix(built):
    import docx
    document = docx.Document(str(built["docx"]))
    tables = [[row.cells[0].text for row in t.rows[1:]] for t in document.tables]
    assert any("Outlet temperature" in rows for rows in tables)
    first = next(rows for rows in tables if "Outlet temperature" in rows)
    assert "Raw Internal Flag A" not in first
    assert any("Raw Internal Flag A" in rows for rows in tables)
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Appendix A. Complete Numerical Results" in text
    html = built["html"]
    assert "-18.5" in html and "Appendix A. Complete Numerical Results" in html


def test_equations_are_editable_with_image_fallback(built):
    xml = _document_xml(built["docx"])
    assert xml.count("<m:oMath") == 1      # Reynolds number
    assert xml.count("<w:drawing") == 2    # figures only; no equation picture is needed
    assert "automatic equation typesetting unavailable" in xml  # \begin{cases} stays marked


def test_discussion_traces_results_without_machine_keys(built):
    html = built["html"]
    assert "Pressure Drop = 3.2 bar; Outlet temperature = -18.5 \u00b0C" in html
    assert "pressure_drop_bar" not in html


def test_html_has_anchors_xrefs_alt_text_and_offline_math(built):
    html = built["html"]
    assert 'id="fig-1"' in html and 'id="fig-2"' in html
    assert 'href="#fig-2"' in html
    assert re.search(r'<img src="data:image/png;base64,[^"]+" alt="Pressure', html)
    assert "cdn.jsdelivr.net" not in html
    assert "KaTeX_Main" in html
    assert '<meta name="author" content="Test Author">' in html


def test_unmarked_long_results_get_a_nudge(tmp_path):
    pytest.importorskip("docx")
    results = {"key_results": {"v{}_kW".format(i): float(i) + 0.5 for i in range(20)},
               "conclusions": "Done."}
    task = _task(tmp_path, results)
    run = subprocess.run([sys.executable, str(task / "step3_report" / "generate_report.py")],
                         cwd=str(task), capture_output=True, text=True)
    assert run.returncode == 0, run.stdout + run.stderr
    assert '"headline": true' in run.stdout
